"""Tests for Excel input handler and schema importer."""

import json
import os
import sys
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, mock_open, patch

import pyarrow as pa
import pytest

from forklift.inputs.config import ExcelInputConfig, ExcelSheetConfig
from forklift.inputs.excel import ExcelInputHandler
from forklift.schema.excel_schema_importer import ExcelSchemaImporter, SchemaValidationError


class TestExcelInputHandler:
    """Test cases for the ExcelInputHandler class."""

    @pytest.fixture
    def basic_config(self):
        """Create a basic Excel input configuration."""
        return ExcelInputConfig(
            sheets=[
                ExcelSheetConfig(select={"name": "employees"}, header={"row": 0}, data_start_row=1)
            ]
        )

    @pytest.fixture
    def multi_sheet_config(self):
        """Create a multi-sheet Excel configuration."""
        return ExcelInputConfig(
            sheets=[
                ExcelSheetConfig(select={"name": "employees"}),
                ExcelSheetConfig(select={"name": "products"}),
                ExcelSheetConfig(select={"name": "sales"}),
            ]
        )

    def test_init(self, basic_config):
        """Test ExcelInputHandler initialization."""
        handler = ExcelInputHandler(basic_config)

        assert handler.config == basic_config
        assert handler._workbook is None
        assert handler._engine is None

    def test_detect_engine_xlsx(self):
        """Test engine detection for .xlsx files."""
        config = ExcelInputConfig()
        handler = ExcelInputHandler(config)

        engine = handler.detect_engine(Path("test.xlsx"))
        assert engine == "openpyxl"

    def test_detect_engine_xls(self):
        """Test engine detection for .xls files."""
        config = ExcelInputConfig()
        handler = ExcelInputHandler(config)

        engine = handler.detect_engine(Path("test.xls"))
        assert engine == "xlrd"

    def test_detect_engine_custom(self):
        """Test using custom engine from config."""
        config = ExcelInputConfig(engine="custom_engine")
        handler = ExcelInputHandler(config)

        engine = handler.detect_engine(Path("test.xlsx"))
        assert engine == "custom_engine"

    def test_detect_engine_unsupported(self):
        """Test error for unsupported file extension."""
        config = ExcelInputConfig()
        handler = ExcelInputHandler(config)

        with pytest.raises(ValueError, match="Unsupported Excel file extension"):
            handler.detect_engine(Path("test.pdf"))

    @patch.object(ExcelInputHandler, "_check_xlsx_archive")
    @patch("openpyxl.load_workbook")
    def test_open_workbook_openpyxl(self, mock_load_workbook, mock_check, basic_config):
        """Test opening workbook with openpyxl engine (read-only streaming mode)."""
        mock_workbook = MagicMock()
        mock_load_workbook.return_value = mock_workbook

        handler = ExcelInputHandler(basic_config)
        handler.open_workbook(Path("test.xlsx"))

        assert handler._engine == "openpyxl"
        assert handler._workbook == mock_workbook
        mock_check.assert_called_once_with(Path("test.xlsx"))
        mock_load_workbook.assert_called_once_with(
            Path("test.xlsx"), read_only=True, data_only=basic_config.values_only
        )

    @patch("xlrd.open_workbook")
    def test_open_workbook_xlrd(self, mock_open_workbook):
        """Test opening workbook with xlrd engine."""
        mock_workbook = MagicMock()
        mock_open_workbook.return_value = mock_workbook

        config = ExcelInputConfig(engine="xlrd")
        handler = ExcelInputHandler(config)
        handler.open_workbook(Path("test.xls"))

        assert handler._engine == "xlrd"
        assert handler._workbook == mock_workbook
        mock_open_workbook.assert_called_once_with("test.xls")

    def test_open_workbook_import_error(self, basic_config):
        """Test ImportError when required library is missing."""
        handler = ExcelInputHandler(basic_config)

        with patch("builtins.__import__", side_effect=ImportError("No module named 'openpyxl'")):
            with pytest.raises(
                ImportError, match="Required library for openpyxl engine not found"
            ):
                handler.open_workbook(Path("test.xlsx"))

    def test_open_workbook_unsupported_engine(self):
        """Test error for unsupported engine."""
        config = ExcelInputConfig(engine="unsupported")
        handler = ExcelInputHandler(config)
        handler._engine = "unsupported"

        with pytest.raises(ValueError, match="Unsupported engine"):
            handler.open_workbook(Path("test.xlsx"))

    def test_close_workbook(self, basic_config):
        """Test closing workbook."""
        handler = ExcelInputHandler(basic_config)
        mock_workbook = MagicMock()
        handler._workbook = mock_workbook

        handler.close_workbook()

        # For openpyxl, close() should be called if available
        # For xlrd, workbook doesn't have close() method

    def test_close_workbook_none(self, basic_config):
        """Test closing workbook when none is open."""
        handler = ExcelInputHandler(basic_config)

        # Should not raise error
        handler.close_workbook()

    @patch.object(ExcelInputHandler, "_check_xlsx_archive")
    @patch("openpyxl.load_workbook")
    def test_get_sheet_names_openpyxl(self, mock_load_workbook, mock_check, basic_config):
        """Test getting sheet names with openpyxl."""
        mock_workbook = MagicMock()
        mock_workbook.sheetnames = ["Sheet1", "Sheet2", "Sheet3"]
        mock_load_workbook.return_value = mock_workbook

        handler = ExcelInputHandler(basic_config)
        handler.open_workbook(Path("test.xlsx"))

        sheet_names = handler.get_sheet_names()
        assert sheet_names == ["Sheet1", "Sheet2", "Sheet3"]

    @patch("xlrd.open_workbook")
    def test_get_sheet_names_xlrd(self, mock_open_workbook):
        """Test getting sheet names with xlrd."""
        mock_workbook = MagicMock()
        mock_workbook.sheet_names.return_value = ["Sheet1", "Sheet2"]
        mock_open_workbook.return_value = mock_workbook

        config = ExcelInputConfig(engine="xlrd")
        handler = ExcelInputHandler(config)
        handler.open_workbook(Path("test.xls"))

        sheet_names = handler.get_sheet_names()
        assert sheet_names == ["Sheet1", "Sheet2"]

    def test_get_sheet_names_no_workbook(self, basic_config):
        """Test getting sheet names when no workbook is open."""
        handler = ExcelInputHandler(basic_config)

        with pytest.raises(
            RuntimeError, match="Workbook not opened. Call open_workbook\\(\\) first."
        ):
            handler.get_sheet_names()

    @staticmethod
    def _write_workbook(path, rows, title="employees"):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title
        for row in rows:
            ws.append(row)
        wb.save(path)
        return path

    def test_read_sheet_data_basic(self, tmp_path):
        """Reading a sheet returns an Arrow table typed column-wise."""
        path = self._write_workbook(tmp_path / "book.xlsx", [["col1", "col2"], [1, 3.5], [2, 4.5]])
        config = ExcelInputConfig(
            sheets=[ExcelSheetConfig(select={"name": "employees"}, header={"row": 0})]
        )

        handler = ExcelInputHandler(config)
        handler.open_workbook(path)
        try:
            result = handler.read_sheet_data("employees", config.sheets[0])
        finally:
            handler.close_workbook()

        assert isinstance(result, pa.Table)
        assert result.to_pydict() == {"col1": [1, 2], "col2": [3.5, 4.5]}
        assert result.schema.field("col1").type == pa.int64()
        assert result.schema.field("col2").type == pa.float64()

    def test_select_sheets_by_name(self, basic_config):
        """Test selecting sheets by name."""
        handler = ExcelInputHandler(basic_config)
        handler._workbook = MagicMock()
        handler._engine = "openpyxl"
        handler._workbook.sheetnames = ["employees", "products", "sales"]

        selected = handler.select_sheets(basic_config.sheets)

        assert len(selected) == 1
        assert selected[0][0] == "employees"

    def test_select_sheets_by_index(self):
        """Test selecting sheets by index."""
        config = ExcelInputConfig(sheets=[ExcelSheetConfig(select={"index": 1})])
        handler = ExcelInputHandler(config)
        handler._workbook = MagicMock()
        handler._engine = "openpyxl"
        handler._workbook.sheetnames = ["Sheet1", "Sheet2", "Sheet3"]

        selected = handler.select_sheets(config.sheets)

        assert len(selected) == 1
        assert selected[0][0] == "Sheet2"  # Index 1

    def test_select_sheets_by_regex(self):
        """Test selecting sheets by regex pattern."""
        config = ExcelInputConfig(sheets=[ExcelSheetConfig(select={"regex": r"Sheet\d+"})])
        handler = ExcelInputHandler(config)
        handler._workbook = MagicMock()
        handler._engine = "openpyxl"
        handler._workbook.sheetnames = ["Sheet1", "Sheet2", "Data", "Sheet3"]

        selected = handler.select_sheets(config.sheets)

        assert len(selected) == 3
        sheet_names = [s[0] for s in selected]
        assert "Sheet1" in sheet_names
        assert "Sheet2" in sheet_names
        assert "Sheet3" in sheet_names
        assert "Data" not in sheet_names

    def test_select_sheets_no_match(self):
        """Test error when no sheets match selection criteria."""
        config = ExcelInputConfig(sheets=[ExcelSheetConfig(select={"name": "nonexistent"})])
        handler = ExcelInputHandler(config)
        handler._workbook = MagicMock()
        handler._engine = "openpyxl"
        handler._workbook.sheetnames = ["Sheet1", "Sheet2"]

        with pytest.raises(ValueError, match="No sheets selected based on configuration criteria"):
            handler.select_sheets(config.sheets)

    def test_read_sheet_data_no_workbook(self, basic_config):
        """Test reading sheet data when no workbook is open."""
        handler = ExcelInputHandler(basic_config)
        sheet_config = basic_config.sheets[0]

        with pytest.raises(
            RuntimeError, match="Workbook not opened. Call open_workbook\\(\\) first."
        ):
            handler.read_sheet_data("employees", sheet_config)

    def test_get_sheet_names_unsupported_engine(self, basic_config):
        """Test getting sheet names with unsupported engine."""
        handler = ExcelInputHandler(basic_config)
        handler._workbook = MagicMock()
        handler._engine = "unsupported_engine"

        with pytest.raises(ValueError, match="Unsupported engine: unsupported_engine"):
            handler.get_sheet_names()

    def test_read_sheet_data_with_na_values(self, tmp_path):
        """na_values (and keep_default_na) turn matching text cells into nulls."""
        path = self._write_workbook(
            tmp_path / "book.xlsx", [["a", "b"], ["N/A", "x"], ["keep", ""], ["NULL", "y"]]
        )
        config = ExcelInputConfig(
            sheets=[ExcelSheetConfig(select={"name": "employees"})],
            na_values=["N/A", "NULL", ""],
            keep_default_na=False,
        )

        handler = ExcelInputHandler(config)
        handler.open_workbook(path)
        try:
            result = handler.read_sheet_data("employees", config.sheets[0])
        finally:
            handler.close_workbook()

        assert result.to_pydict() == {"a": [None, "keep", None], "b": ["x", None, "y"]}

    def test_read_sheet_data_with_data_end_row(self, tmp_path):
        """data_start_row/data_end_row are inclusive 1-based sheet rows."""
        path = self._write_workbook(
            tmp_path / "book.xlsx", [["n"]] + [[i] for i in range(1, 13)]  # data on rows 2..13
        )
        config = ExcelInputConfig(
            sheets=[
                ExcelSheetConfig(select={"name": "employees"}, data_start_row=2, data_end_row=10)
            ]
        )

        handler = ExcelInputHandler(config)
        handler.open_workbook(path)
        try:
            result = handler.read_sheet_data("employees", config.sheets[0])
        finally:
            handler.close_workbook()

        assert result["n"].to_pylist() == list(range(1, 10))  # sheet rows 2..10 -> 9 rows

    def test_select_sheets_no_workbook(self, basic_config):
        """Test selecting sheets when no workbook is open."""
        handler = ExcelInputHandler(basic_config)

        with pytest.raises(
            RuntimeError, match="Workbook not opened. Call open_workbook\\(\\) first."
        ):
            handler.select_sheets(basic_config.sheets)
