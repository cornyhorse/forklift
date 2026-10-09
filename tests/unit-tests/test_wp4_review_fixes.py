"""Regression tests for the WP4 review fixes (inputs/, io/, encoding detection).

Covers: Excel input (previously missing ``get_sheet_info``/``process_sheets``), SQL identifier
safety and type mapping, S3 writer/path/IO handling, FWF parsing and validation, encoding
detection and the CSV header search. Everything runs without a database or network: pyodbc is
replaced by a fake module and S3 by moto.
"""

from __future__ import annotations

import builtins
import datetime
import decimal
import io
import shutil
import sys
import types
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from forklift.inputs.config import (
    CsvInputConfig,
    ExcelInputConfig,
    ExcelSheetConfig,
    FwfConditionalSchema,
    FwfFieldSpec,
    FwfInputConfig,
    SqlInputConfig,
)
from forklift.inputs.csv import CsvInputHandler
from forklift.inputs.excel import ExcelInputHandler
from forklift.inputs.fwf import FwfInputHandler, FwfTypeConverter
from forklift.inputs.sql import SqlInputHandler
from forklift.inputs.sql.connection import SqlConnectionManager
from forklift.inputs.sql.types import SqlTypeConverter
from forklift.io.s3_streaming import (
    S3Path,
    S3StreamingClient,
    S3StreamingWriter,
    is_s3_path,
    normalize_s3_uri,
)
from forklift.io.unified_io import S3ParquetWriter, UnifiedCSVWriter, UnifiedIOHandler
from forklift.utils.detect_encoding import detect_encoding, open_text_auto, verify_encoding

TEST_DATA = Path(__file__).parent.parent / "test-data"


# =====================================================================================
# 1. Excel input
# =====================================================================================


def _write_xlsx(path: Path, rows, title: str = "Sheet1", extra_sheets=None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    for name, sheet_rows in (extra_sheets or {}).items():
        extra = wb.create_sheet(name)
        for row in sheet_rows:
            extra.append(row)
    wb.save(path)
    return path


def _read(path: Path, **sheet_kwargs) -> pa.Table:
    config = ExcelInputConfig(sheets=[ExcelSheetConfig(select={"index": 0}, **sheet_kwargs)])
    ((name, table),) = list(ExcelInputHandler(config).process_sheets(path))
    return table


class TestExcelHandlerApi:
    def test_get_sheet_info_and_process_sheets_exist_on_repo_file(self):
        """The importer's calls used to raise AttributeError for every workbook."""
        handler = ExcelInputHandler(ExcelInputConfig())
        info = handler.get_sheet_info(TEST_DATA / "multi_sheet.xlsx")
        assert info["engine"] == "openpyxl"
        assert info["sheet_count"] == 3
        assert info["sheet_names"] == ["employees", "products", "sales"]
        assert handler._workbook is None  # nothing left open

        results = list(handler.process_sheets(TEST_DATA / "multi_sheet.xlsx"))
        assert [name for name, _ in results] == ["employees", "products", "sales"]
        employees = results[0][1]
        assert isinstance(employees, pa.Table)
        assert employees.num_rows == 5
        assert employees.column_names == ["id", "name", "age", "salary", "active"]
        assert employees.schema.field("active").type == pa.bool_()
        assert results[2][1].schema.field("sale_date").type == pa.timestamp("us")

    def test_import_excel_end_to_end(self, tmp_path):
        import forklift as fl

        results = fl.import_excel(str(TEST_DATA / "simple.xlsx"), str(tmp_path / "out"))
        assert results.total_rows == 5
        (output,) = results.output_files
        assert pq.read_table(output).column_names == ["id", "name", "age", "salary", "active"]

    def test_no_pandas_in_excel_module(self):
        source = (Path(__file__).parents[2] / "src/forklift/inputs/excel.py").read_text()
        assert "pandas" not in source.replace("pandas-based", "")
        assert "import pandas" not in source

    def test_workbook_is_closed_when_the_generator_is_abandoned(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", [["a"], [1]])
        handler = ExcelInputHandler(ExcelInputConfig())
        generator = handler.process_sheets(path)
        next(generator)
        assert handler._workbook is not None
        generator.close()
        assert handler._workbook is None


class TestExcelRowSemantics:
    ROWS = [["id", "name"]] + [[i, f"n{i}"] for i in range(1, 7)]  # header row 1, data rows 2..7

    def test_default_header_is_row_one_and_no_row_is_lost(self, tmp_path):
        table = _read(_write_xlsx(tmp_path / "a.xlsx", self.ROWS))
        assert table.column_names == ["id", "name"]
        assert table["id"].to_pylist() == [1, 2, 3, 4, 5, 6]

    def test_header_row_zero_and_data_start_row_two(self, tmp_path):
        """Old code skipped the header and dropped the first data row."""
        path = _write_xlsx(tmp_path / "a.xlsx", self.ROWS)
        table = _read(path, header={"row": 0}, data_start_row=2)
        assert table.column_names == ["id", "name"]
        assert table["id"].to_pylist() == [1, 2, 3, 4, 5, 6]

    def test_data_end_row_is_inclusive_sheet_row(self, tmp_path):
        """data_end_row=5 means sheet rows 2..5, i.e. four data rows (old: five)."""
        path = _write_xlsx(tmp_path / "a.xlsx", self.ROWS)
        table = _read(path, header={"row": 0}, data_start_row=2, data_end_row=5)
        assert table["id"].to_pylist() == [1, 2, 3, 4]

    def test_data_start_row_skips_rows(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", self.ROWS)
        table = _read(path, data_start_row=5)
        assert table["id"].to_pylist() == [4, 5, 6]

    def test_data_start_never_precedes_the_header(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", self.ROWS)
        table = _read(path, header={"row": 0}, data_start_row=1)
        assert table["id"].to_pylist() == [1, 2, 3, 4, 5, 6]

    def test_header_below_metadata_rows(self, tmp_path):
        rows = [["report"], [None], ["id", "name"], [1, "a"], [2, "b"]]
        table = _read(_write_xlsx(tmp_path / "a.xlsx", rows), header={"row": 2})
        assert table.column_names == ["id", "name"]
        assert table["name"].to_pylist() == ["a", "b"]

    def test_invalid_row_settings_raise(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", self.ROWS)
        with pytest.raises(ValueError, match="data_start_row"):
            _read(path, data_start_row=0)
        with pytest.raises(ValueError, match="data_end_row"):
            _read(path, data_start_row=5, data_end_row=3)
        with pytest.raises(ValueError, match="header mode"):
            _read(path, header={"mode": "sometimes"})
        with pytest.raises(ValueError, match="header.row"):
            _read(path, header={"row": -1})


class TestExcelHeaders:
    def test_absent_mode_generates_col_names_and_keeps_first_row(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", [[1, "a"], [2, "b"]])
        table = _read(path, header={"mode": "absent"})
        assert table.column_names == ["col_1", "col_2"]
        assert table["col_1"].to_pylist() == [1, 2]

    def test_absent_mode_with_override_names(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", [[1, "a"], [2, "b"]])
        table = _read(path, header={"mode": "absent", "override": ["id", "label"]})
        assert table.column_names == ["id", "label"]
        assert table.num_rows == 2

    def test_auto_mode_detects_text_header_or_falls_back_to_col_names(self, tmp_path):
        with_header = _write_xlsx(tmp_path / "h.xlsx", [["id", "name"], [1, "a"]])
        assert _read(with_header, header={"mode": "auto"}).column_names == ["id", "name"]

        without = _write_xlsx(tmp_path / "n.xlsx", [[1, "a"], [2, "b"]])
        table = _read(without, header={"mode": "auto"})
        assert table.column_names == ["col_1", "col_2"]
        assert table.num_rows == 2

    def test_blank_and_duplicate_header_names_are_made_unique(self, tmp_path):
        """Three or more blank headers used to be an infinite-loop risk in the deduper."""
        rows = [["id", None, None, None, "id", "id"], [1, 2, 3, 4, 5, 6]]
        table = _read(_write_xlsx(tmp_path / "a.xlsx", rows))
        assert table.column_names == ["id", "col_2", "col_3", "col_4", "id_1", "id_2"]

    def test_sheet_with_repeated_header_names_from_the_repo_file(self):
        config = ExcelInputConfig(
            sheets=[ExcelSheetConfig(select={"name": "Sheet3"}, header={"row": 4})]
        )
        path = Path(__file__).parent.parent / "test-files" / "excel" / "excel-data.xlsx"
        ((_, table),) = list(ExcelInputHandler(config).process_sheets(path))
        assert table.column_names[:4] == ["id", "name", "name_1", "amount_usd"]


class TestExcelSettings:
    def test_nulls_na_values_and_per_column_nulls(self, tmp_path):
        rows = [["a", "b"], ["N/A", "-"], ["x", "keep"], ["  ", "z"]]
        path = _write_xlsx(tmp_path / "a.xlsx", rows)

        config = ExcelInputConfig(
            sheets=[ExcelSheetConfig(select={"index": 0})],
            nulls={"global": ["N/A"], "perColumn": {"b": ["-"]}},
        )
        ((_, table),) = list(ExcelInputHandler(config).process_sheets(path))
        assert table["a"].to_pylist() == [None, "x", None]
        assert table["b"].to_pylist() == [None, "keep", "z"]

        config = ExcelInputConfig(
            sheets=[ExcelSheetConfig(select={"index": 0})], na_values=["x"], keep_default_na=False
        )
        ((_, table),) = list(ExcelInputHandler(config).process_sheets(path))
        assert table["a"].to_pylist() == ["N/A", None, "  "]

    def test_numeric_na_values(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", [["n"], [1], [-999], [3]])
        config = ExcelInputConfig(
            sheets=[ExcelSheetConfig(select={"index": 0})], na_values=["-999"]
        )
        ((_, table),) = list(ExcelInputHandler(config).process_sheets(path))
        assert table["n"].to_pylist() == [1, None, 3]

    def test_skip_blank_rows(self, tmp_path):
        rows = [["a", "b"], [1, 2], [None, None], [3, 4], [None, None]]
        path = _write_xlsx(tmp_path / "a.xlsx", rows)
        assert _read(path).num_rows == 2
        kept = _read(path, skip_blank_rows=False)
        assert kept["a"].to_pylist() == [1, None, 3]  # interior blank kept, trailing trimmed

    def test_columns_select_rename_and_cast(self, tmp_path):
        rows = [["id", "name", "price"], [1, "a", "1.50"], [2, "b", "2.25"]]
        path = _write_xlsx(tmp_path / "a.xlsx", rows)
        table = _read(
            path,
            columns=[
                {"name": "price_usd", "position": "C", "parquetType": "double"},
                {"name": "key", "position": 1, "parquetType": "string"},
            ],
        )
        assert table.column_names == ["price_usd", "key"]
        assert table["price_usd"].to_pylist() == [1.5, 2.25]
        assert table["key"].to_pylist() == ["1", "2"]

        by_name = _read(path, columns=["name", "id"])
        assert by_name.column_names == ["name", "id"]

    def test_columns_errors_do_not_leak_values(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", [["v"], ["secret-text"]])
        with pytest.raises(ValueError, match="cannot be converted") as excinfo:
            _read(path, columns=[{"name": "v", "position": "A", "parquetType": "int64"}])
        assert "secret-text" not in str(excinfo.value)
        with pytest.raises(ValueError, match="not found"):
            _read(path, columns=["missing"])
        with pytest.raises(ValueError, match="Unsupported parquetType"):
            _read(path, columns=[{"name": "v", "position": "A", "parquetType": "weird"}])

    def test_name_override_names_the_output(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", [["a"], [1]])
        config = ExcelInputConfig(
            sheets=[ExcelSheetConfig(select={"index": 0}, name_override="renamed")]
        )
        ((name, _),) = list(ExcelInputHandler(config).process_sheets(path))
        assert name == "renamed"

    def test_mixed_type_columns_become_strings(self, tmp_path):
        rows = [
            ["m", "i", "f", "d"],
            [1, 1, 1.5, datetime.datetime(2024, 1, 2)],
            ["x", 2, 2, None],
        ]
        table = _read(_write_xlsx(tmp_path / "a.xlsx", rows))
        assert table.schema.field("m").type == pa.string()
        assert table["m"].to_pylist() == ["1", "x"]
        assert table.schema.field("i").type == pa.int64()
        assert table.schema.field("f").type == pa.float64()
        assert table.schema.field("d").type == pa.timestamp("us")

    def test_regex_is_compiled_at_config_time(self):
        with pytest.raises(ValueError, match="select.regex"):
            ExcelSheetConfig(select={"regex": "("})

    def test_select_sheets_raises_for_missing_name_or_index(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", [["a"], [1]], extra_sheets={"Two": [["b"], [2]]})
        for select in ({"name": "Nope"}, {"index": 5}, {"index": -1}, {"regex": "^zzz"}, {}):
            config = ExcelInputConfig(sheets=[ExcelSheetConfig(select=select)])
            with pytest.raises(ValueError, match="No sheets selected"):
                list(ExcelInputHandler(config).process_sheets(path))

    def test_sheets_default_to_all_when_unset(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", [["a"], [1]], extra_sheets={"Two": [["b"], [2]]})
        names = [n for n, _ in ExcelInputHandler(ExcelInputConfig()).process_sheets(path)]
        assert names == ["Sheet1", "Two"]


class TestExcelResourceLimits:
    def test_sparse_sheet_is_not_expanded_to_its_bounding_box(self, tmp_path):
        """One stray cell at row 100000 / column 30 took 11 s and 839 MB through pandas."""
        import time
        import tracemalloc

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"], ws["B1"], ws["A2"] = "a", "b", 1
        ws.cell(row=100000, column=30, value="x")
        path = tmp_path / "sparse.xlsx"
        wb.save(path)

        tracemalloc.start()
        started = time.time()
        table = _read(path)
        elapsed = time.time() - started
        peak = tracemalloc.get_traced_memory()[1]
        tracemalloc.stop()

        assert table.num_rows == 2  # row 2 and the stray row; blank rows in between skipped
        assert table.num_columns == 30
        assert elapsed < 5
        assert peak < 20 * 1024 * 1024

    def test_trailing_empty_rows_and_columns_are_trimmed(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["a", "b", None, None])
        ws.append([1, 2, None, None])
        ws.append([None] * 4)
        ws.cell(row=10, column=8).value = None  # extends the used range without content
        path = tmp_path / "a.xlsx"
        wb.save(path)
        table = _read(path)
        assert table.column_names == ["a", "b"]
        assert table.num_rows == 1

    def test_max_rows_and_max_cells_raise_clear_errors(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", [["a", "b"]] + [[i, i] for i in range(50)])
        for limits in ({"max_rows": 10}, {"max_cells": 20}):
            config = ExcelInputConfig(sheets=[ExcelSheetConfig(select={"index": 0})], **limits)
            with pytest.raises(ValueError, match="max_rows|max_cells"):
                list(ExcelInputHandler(config).process_sheets(path))

    def test_zip_bomb_is_refused_before_opening(self, tmp_path):
        path = tmp_path / "bomb.xlsx"
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("xl/worksheets/sheet1.xml", b"\x00" * (40 * 1024 * 1024))

        handler = ExcelInputHandler(ExcelInputConfig())
        with patch("openpyxl.load_workbook") as load_workbook:
            with pytest.raises(ValueError, match="compression ratio"):
                handler.get_sheet_info(path)
            load_workbook.assert_not_called()

    def test_uncompressed_size_limit(self, tmp_path):
        path = _write_xlsx(tmp_path / "a.xlsx", [["a"], [1]])
        handler = ExcelInputHandler(ExcelInputConfig(max_uncompressed_bytes=100))
        with pytest.raises(ValueError, match="max_uncompressed_bytes"):
            handler.get_sheet_info(path)

    def test_non_zip_file_is_a_clear_error(self, tmp_path):
        path = tmp_path / "fake.xlsx"
        path.write_bytes(b"not a zip")
        with pytest.raises(ValueError, match="not a valid .xlsx"):
            ExcelInputHandler(ExcelInputConfig()).get_sheet_info(path)


class TestExcelXls:
    def test_xlrd_rows_are_converted(self):
        import xlrd

        cells = [
            [
                xlrd.sheet.Cell(xlrd.XL_CELL_TEXT, "id"),
                xlrd.sheet.Cell(xlrd.XL_CELL_TEXT, "when"),
                xlrd.sheet.Cell(xlrd.XL_CELL_TEXT, "ok"),
            ],
            [
                xlrd.sheet.Cell(xlrd.XL_CELL_NUMBER, 7.0),
                xlrd.sheet.Cell(xlrd.XL_CELL_DATE, 45292.0),
                xlrd.sheet.Cell(xlrd.XL_CELL_BOOLEAN, 1),
            ],
            [
                xlrd.sheet.Cell(xlrd.XL_CELL_NUMBER, 8.5),
                xlrd.sheet.Cell(xlrd.XL_CELL_EMPTY, ""),
                xlrd.sheet.Cell(xlrd.XL_CELL_BOOLEAN, 0),
            ],
        ]
        sheet = SimpleNamespace(nrows=3, ncols=3, row=lambda i: cells[i])
        workbook = SimpleNamespace(
            datemode=0,
            sheet_names=lambda: ["S"],
            sheet_by_name=lambda name: sheet,
            release_resources=lambda: None,
        )

        handler = ExcelInputHandler(
            ExcelInputConfig(sheets=[ExcelSheetConfig(select={"index": 0})])
        )
        handler._workbook, handler._engine = workbook, "xlrd"
        table = handler.read_sheet_data("S", handler.config.sheets[0])
        assert table.column_names == ["id", "when", "ok"]
        assert table["id"].to_pylist() == [7, 8.5]
        assert table["when"].to_pylist() == [datetime.datetime(2024, 1, 1), None]
        assert table["ok"].to_pylist() == [True, False]


# =====================================================================================
# 2./3. SQL
# =====================================================================================


@pytest.fixture
def fake_pyodbc(monkeypatch):
    module = types.ModuleType("pyodbc")
    for name, value in dict(
        SQL_CHAR=1,
        SQL_NUMERIC=2,
        SQL_DECIMAL=3,
        SQL_INTEGER=4,
        SQL_SMALLINT=5,
        SQL_FLOAT=6,
        SQL_REAL=7,
        SQL_DOUBLE=8,
        SQL_VARCHAR=12,
        SQL_TYPE_DATE=91,
        SQL_TYPE_TIME=92,
        SQL_TYPE_TIMESTAMP=93,
        SQL_LONGVARCHAR=-1,
        SQL_BINARY=-2,
        SQL_VARBINARY=-3,
        SQL_LONGVARBINARY=-4,
        SQL_BIGINT=-5,
        SQL_TINYINT=-6,
        SQL_BIT=-7,
        SQL_WCHAR=-8,
        SQL_WVARCHAR=-9,
        SQL_WLONGVARCHAR=-10,
        SQL_IDENTIFIER_QUOTE_CHAR=29,
    ).items():
        setattr(module, name, value)
    module.pooling = True
    module.connect = MagicMock(name="pyodbc.connect")
    monkeypatch.setitem(sys.modules, "pyodbc", module)
    return module


class FakeCursor:
    """Just enough of a pyodbc cursor, recording every statement it is asked to run."""

    def __init__(self, catalog, columns=None, rows=None, description=None, columns_error=None):
        self.catalog = catalog  # [(schema or None, table, type)]
        self.column_rows = columns or []
        self.rows = list(rows or [])
        self.description = description
        self.columns_error = columns_error
        self.executed = []
        self.columns_kwargs = []

    def tables(self, **kwargs):
        return [
            SimpleNamespace(table_schem=s, table_name=t, table_type=kind)
            for s, t, kind in self.catalog
        ]

    def columns(self, **kwargs):
        self.columns_kwargs.append(kwargs)
        if self.columns_error:
            raise self.columns_error
        return list(self.column_rows)

    def execute(self, sql):
        self.executed.append(sql)

    def fetchmany(self, size):
        batch, self.rows = self.rows[:size], self.rows[size:]
        return batch

    def close(self):
        pass


class FakeConnection:
    def __init__(self, cursor, quote='"'):
        self._cursor = cursor
        self.quote = quote

    def cursor(self):
        return self._cursor

    def getinfo(self, what):
        return self.quote


def _column(table, name, type_name, schema="main", size=None, digits=None):
    return SimpleNamespace(
        table_schem=schema,
        table_name=table,
        column_name=name,
        type_name=type_name,
        column_size=size,
        decimal_digits=digits,
        nullable=True,
    )


def _sql_handler(cursor, quote='"', **config_kwargs):
    handler = SqlInputHandler(SqlInputConfig(connection_string="DSN=x", **config_kwargs))
    handler.connection = FakeConnection(cursor, quote)
    return handler


class TestSqlIdentifierSafety:
    def test_injection_via_table_name_is_rejected_before_any_sql_runs(self, fake_pyodbc):
        cursor = FakeCursor(
            [("main", "employees", "TABLE")], [_column("employees", "id", "INTEGER")]
        )
        handler = _sql_handler(cursor)
        evil = 'employees" WHERE 1=0 UNION SELECT password FROM users --'

        with pytest.raises(ValueError, match="not found in the database catalog"):
            list(handler.read_table_data("main", evil))
        with pytest.raises(ValueError, match="not found in the database catalog"):
            list(handler.read_table_data('main" OR 1=1 --', "employees"))
        with pytest.raises(ValueError):
            handler.get_table_schema("main", evil)
        assert cursor.executed == []

    def test_identifiers_are_always_quoted_and_quotes_are_doubled(self, fake_pyodbc):
        odd = 'we"ird'
        cursor = FakeCursor(
            [("sch", odd, "TABLE")],
            [_column(odd, "id", "INTEGER", schema="sch")],
            rows=[(1,)],
        )
        handler = _sql_handler(cursor)  # use_quoted_identifiers defaults to False

        batches = list(handler.read_table_data("sch", odd))

        assert cursor.executed == ['SELECT * FROM "sch"."we""ird"']
        assert batches[0].to_pydict() == {"id": [1]}

    def test_quote_character_comes_from_the_driver(self, fake_pyodbc):
        cursor = FakeCursor([("db", "t", "TABLE")], [_column("t", "id", "INTEGER", schema="db")])
        handler = _sql_handler(cursor, quote="`")
        assert handler._quote_identifier("a`b") == "`a``b`"
        list(handler.read_table_data("db", "t"))
        assert cursor.executed == ["SELECT * FROM `db`.`t`"]

    def test_quote_falls_back_to_double_quote(self, fake_pyodbc):
        handler = _sql_handler(FakeCursor([]), quote=" ")  # ODBC: space = no quote char
        assert handler._quote_identifier("t") == '"t"'
        with pytest.raises(ValueError):
            handler._quote_identifier("")
        with pytest.raises(ValueError):
            handler._quote_identifier("a\x00b")

    def test_unknown_table_raises_value_error(self, fake_pyodbc):
        handler = _sql_handler(FakeCursor([("main", "employees", "TABLE")]))
        with pytest.raises(ValueError, match="not found in the database catalog"):
            handler.get_table_schema("main", "no_such_table")
        with pytest.raises(ValueError, match="not found in the database catalog"):
            handler.get_table_schema("other", "employees")

    def test_case_differences_resolve_to_the_catalog_spelling_when_unambiguous(self, fake_pyodbc):
        cursor = FakeCursor(
            [("Main", "Employees", "TABLE")],
            [_column("Employees", "id", "INTEGER", schema="Main")],
            rows=[(1,)],
        )
        handler = _sql_handler(cursor)
        assert handler.schema_manager.resolve_table("main", "EMPLOYEES") == ("Main", "Employees")
        list(handler.read_table_data("main", "EMPLOYEES"))
        assert cursor.executed == ['SELECT * FROM "Main"."Employees"']

        ambiguous = _sql_handler(FakeCursor([("s", "Users", "TABLE"), ("s", "users", "TABLE")]))
        with pytest.raises(ValueError, match="ambiguous"):
            ambiguous.schema_manager.resolve_table("s", "USERS")
        # an exact spelling still wins over the case-insensitive candidates
        assert ambiguous.schema_manager.resolve_table("s", "users") == ("s", "users")

    def test_columns_are_filtered_to_the_exact_table(self, fake_pyodbc):
        """my_table matches myXtable as an ODBC search pattern."""
        cursor = FakeCursor(
            [("main", "my_table", "TABLE")],
            [
                _column("my_table", "id", "INTEGER"),
                _column("myXtable", "intruder", "VARCHAR"),
                _column("my_table", "name", "VARCHAR"),
            ],
        )
        schema = _sql_handler(cursor).get_table_schema("main", "my_table")
        assert schema.names == ["id", "name"]

    def test_empty_columns_result_raises_instead_of_zero_column_schema(self, fake_pyodbc):
        cursor = FakeCursor([("main", "t", "TABLE")], [])
        with pytest.raises(ValueError, match="no columns"):
            _sql_handler(cursor).get_table_schema("main", "t")

    def test_fallback_uses_where_1_equals_0_and_python_types(self, fake_pyodbc):
        cursor = FakeCursor(
            [(None, "t", "TABLE")],
            columns_error=RuntimeError("columns() unsupported"),
            description=[
                ("i", int, None, None, None, None, True),
                ("s", str, None, None, None, None, True),
                ("f", float, None, None, None, None, True),
                ("d", decimal.Decimal, None, None, 12, 3, True),
                ("ts", datetime.datetime, None, None, None, None, False),
                ("dt", datetime.date, None, None, None, None, True),
                ("b", bool, None, None, None, None, True),
                ("raw", bytes, None, None, None, None, True),
            ],
        )
        schema = _sql_handler(cursor).get_table_schema("default", "t")

        assert cursor.executed == ['SELECT * FROM "t" WHERE 1=0']
        assert [f.type for f in schema] == [
            pa.int64(),
            pa.string(),
            pa.float64(),
            pa.decimal128(12, 3),
            pa.timestamp("us"),
            pa.date32(),
            pa.bool_(),
            pa.binary(),
        ]
        assert schema.field("ts").nullable is False

    def test_fallback_does_not_run_for_unknown_table(self, fake_pyodbc):
        cursor = FakeCursor([("main", "t", "TABLE")], columns_error=RuntimeError("unsupported"))
        with pytest.raises(ValueError, match="not found"):
            _sql_handler(cursor).get_table_schema("main", 'x"; DROP TABLE t; --')
        assert cursor.executed == []

    def test_catalog_failure_raises_instead_of_returning_no_tables(self, fake_pyodbc):
        cursor = FakeCursor([])
        cursor.tables = MagicMock(side_effect=RuntimeError("no catalog"))
        cursor.execute = MagicMock(side_effect=RuntimeError("no sqlite_master"))
        with pytest.raises(RuntimeError, match="Could not retrieve the table list"):
            _sql_handler(cursor).get_table_list()

    def test_bare_table_name_in_several_schemas_is_ambiguous(self, fake_pyodbc):
        cursor = FakeCursor([("a", "users", "TABLE"), ("b", "users", "TABLE")])
        handler = _sql_handler(cursor)
        with pytest.raises(ValueError, match="ambiguous"):
            handler.get_specified_tables(["users"])
        with pytest.raises(ValueError, match="ambiguous"):
            handler.schema_manager.resolve_table("default", "users")
        assert handler.get_specified_tables(["b.users"]) == [("b", "users")]

    def test_explicit_schema_does_not_fall_back_to_another_schema(self, fake_pyodbc):
        handler = _sql_handler(FakeCursor([("a", "users", "TABLE")]))
        assert handler.get_specified_tables(["zzz.users"]) == []

    def test_schemaless_catalog_accepts_a_schema_qualified_spec(self, fake_pyodbc):
        cursor = FakeCursor([(None, "users", "TABLE")], [_column("users", "id", "INTEGER", None)])
        handler = _sql_handler(cursor)
        assert handler.schema_manager.resolve_table("main", "users") == (None, "users")


class TestSqlConnection:
    def test_connect_is_read_only_by_default(self, fake_pyodbc):
        manager = SqlConnectionManager(SqlInputConfig(connection_string="DSN=x"))
        manager.connect()
        fake_pyodbc.connect.assert_called_once_with("DSN=x", timeout=30, readonly=True)

    def test_read_only_can_be_disabled(self, fake_pyodbc):
        manager = SqlConnectionManager(SqlInputConfig(connection_string="DSN=x", read_only=False))
        manager.connect()
        fake_pyodbc.connect.assert_called_once_with("DSN=x", timeout=30)

    def test_connection_params_are_brace_escaped(self, fake_pyodbc):
        params = {"PWD": "p;ass}word", "Plain": "ok", "Empty": "", "Eq": "a=b"}
        manager = SqlConnectionManager(
            SqlInputConfig(connection_string="DSN=x", connection_params=params)
        )
        manager.connect()
        (conn_str,), _ = fake_pyodbc.connect.call_args
        assert conn_str == "DSN=x;PWD={p;ass}}word};Plain=ok;Empty={};Eq={a=b}"

    def test_connection_param_names_cannot_inject(self, fake_pyodbc):
        manager = SqlConnectionManager(
            SqlInputConfig(connection_string="DSN=x", connection_params={"A=1;UID": "x"})
        )
        with pytest.raises(ValueError, match="parameter name"):
            manager.connect()
        fake_pyodbc.connect.assert_not_called()

    def test_repr_does_not_leak_credentials(self):
        config = SqlInputConfig(
            connection_string="DRIVER={x};UID=u;PWD=hunter2",
            connection_params={"PWD": "hunter3"},
        )
        text = repr(config)
        assert "hunter2" not in text and "hunter3" not in text
        assert "batch_size=10000" in text
        assert config.connection_string.endswith("hunter2")  # still usable


class TestSqlTypes:
    CONVERTER = SqlTypeConverter()

    @pytest.mark.parametrize(
        "sql_type,size,digits,expected",
        [
            ("FLOAT", None, None, pa.float64()),
            ("float", None, None, pa.float64()),
            ("REAL", None, None, pa.float32()),
            ("DOUBLE", None, None, pa.float64()),
            ("int identity", None, None, pa.int32()),
            ("bigint identity", None, None, pa.int64()),
            ("smallint identity", None, None, pa.int16()),
            ("datetime2", None, None, pa.timestamp("us")),
            ("smalldatetime", None, None, pa.timestamp("us")),
            ("money", None, None, pa.decimal128(19, 4)),
            ("smallmoney", None, None, pa.decimal128(10, 4)),
            ("uniqueidentifier", None, None, pa.string()),
            ("SMALLINT", None, None, pa.int16()),
            ("TIME", None, None, pa.time64("us")),
            ("DECIMAL", 18, 4, pa.decimal128(18, 4)),
            ("numeric", 38, 10, pa.decimal128(38, 10)),
            ("DECIMAL", 40, 2, pa.string()),
            ("NUMERIC", 76, 10, pa.string()),
            ("DECIMAL", 10, 12, pa.string()),
        ],
    )
    def test_type_mapping(self, sql_type, size, digits, expected):
        assert self.CONVERTER.sql_type_to_pyarrow(sql_type, size, digits) == expected

    def test_float_no_longer_loses_precision(self):
        array = self.CONVERTER.convert_column_data(
            (123456789.123,), self.CONVERTER.sql_type_to_pyarrow("FLOAT")
        )
        assert array.to_pylist() == [123456789.123]

    def test_wide_decimal_column_does_not_crash_the_table(self, fake_pyodbc):
        cursor = FakeCursor(
            [("main", "t", "TABLE")],
            [
                _column("t", "id", "INTEGER"),
                _column("t", "huge", "DECIMAL", size=50, digits=2),
            ],
            rows=[(1, decimal.Decimal("123456789012345678901234567890123456789012345.67"))],
        )
        (batch,) = list(_sql_handler(cursor).read_table_data("main", "t"))
        assert batch.schema.field("huge").type == pa.string()
        assert batch["huge"].to_pylist() == ["123456789012345678901234567890123456789012345.67"]

    def test_string_fallback_never_contradicts_the_declared_type(self):
        with pytest.raises(ValueError, match="declared type int32") as excinfo:
            self.CONVERTER.convert_column_data(("abc-secret",), pa.int32())
        assert "abc-secret" not in str(excinfo.value)

        # A cast the driver's values allow is still performed (ISO strings -> timestamp)
        array = self.CONVERTER.convert_column_data(("2024-01-02 03:04:05",), pa.timestamp("us"))
        assert array.type == pa.timestamp("us")
        # Text columns may hold anything
        assert self.CONVERTER.convert_column_data((1, "x"), pa.string()).to_pylist() == ["1", "x"]

    def test_python_type_mapping(self):
        mapping = {
            bool: "BOOLEAN",
            int: "BIGINT",
            float: "DOUBLE",
            decimal.Decimal: "DECIMAL",
            datetime.datetime: "TIMESTAMP",
            datetime.date: "DATE",
            datetime.time: "TIME",
            bytes: "VARBINARY",
            str: "VARCHAR",
        }
        for python_type, expected in mapping.items():
            assert self.CONVERTER.python_type_to_string(python_type) == expected
        assert self.CONVERTER.python_type_to_string(object) == "VARCHAR"


# =====================================================================================
# 4. S3 / unified I/O
# =====================================================================================


@pytest.fixture
def s3(monkeypatch):
    pytest.importorskip("moto")  # dev-only dependency; S3 tests are skipped without it
    from moto import mock_aws

    for name in ("AWS_ACCESS_KEY_ID", "AWS_SECRET_ACCESS_KEY", "AWS_SESSION_TOKEN"):
        monkeypatch.setenv(name, "testing")
    monkeypatch.setenv("AWS_DEFAULT_REGION", "us-east-1")
    with mock_aws():
        client = S3StreamingClient(
            aws_access_key_id="testing", aws_secret_access_key="testing", region_name="us-east-1"
        )
        client._s3_client.create_bucket(Bucket="bkt")
        yield client


def _object_body(client: S3StreamingClient, key: str) -> bytes:
    return client._s3_client.get_object(Bucket="bkt", Key=key)["Body"].read()


def _open_uploads(client: S3StreamingClient) -> list:
    return client._s3_client.list_multipart_uploads(Bucket="bkt").get("Uploads", [])


class TestS3Writer:
    def test_exception_in_with_body_does_not_publish_a_truncated_object(self, s3):
        with pytest.raises(RuntimeError, match="boom"):
            with s3.open_for_write("s3://bkt/out.json") as writer:
                writer.write('{"a": [1, 2,')
                raise RuntimeError("boom")

        assert not s3.exists("s3://bkt/out.json")
        assert _open_uploads(s3) == []  # the multipart upload was aborted, not left dangling

    def test_clean_exit_publishes_the_object(self, s3):
        with s3.open_for_write("s3://bkt/out.json") as writer:
            writer.write('{"a": [1, 2]}')
        assert _object_body(s3, "out.json") == b'{"a": [1, 2]}'

    def test_abort_is_idempotent_and_blocks_later_publishing(self, s3):
        writer = s3.open_for_write("s3://bkt/never.txt")
        writer.write("partial")
        writer.abort()
        writer.abort()
        writer.close()  # no-op after abort
        assert writer.closed
        assert not s3.exists("s3://bkt/never.txt")
        with pytest.raises(ValueError):
            writer.write("more")

    def test_abort_after_close_keeps_the_published_object(self, s3):
        writer = s3.open_for_write("s3://bkt/kept.txt")
        writer.write("data")
        writer.close()
        writer.abort()
        assert _object_body(s3, "kept.txt") == b"data"

    def test_zero_byte_write_creates_an_empty_object(self, s3):
        with s3.open_for_write("s3://bkt/empty.txt"):
            pass
        assert s3.exists("s3://bkt/empty.txt")
        assert s3.get_size("s3://bkt/empty.txt") == 0

    def test_tell_counts_bytes_and_bytes_like_data_is_accepted(self, s3):
        writer = s3.open_for_write("s3://bkt/mixed.bin")
        assert writer.write("é") == 1  # one character...
        assert writer.tell() == 2  # ...two bytes
        assert writer.write(bytearray(b"ab")) == 2
        assert writer.write(memoryview(b"cde")) == 3
        assert writer.tell() == 7
        with pytest.raises(TypeError):
            writer.write(123)
        writer.close()
        assert _object_body(s3, "mixed.bin") == "é".encode() + b"abcde"

    def test_multipart_path_still_works(self, s3):
        writer = S3StreamingWriter(
            s3._s3_client, S3Path("s3://bkt/big.bin"), part_size=5 * 1024 * 1024, mode="wb"
        )
        chunk = b"x" * (3 * 1024 * 1024)
        for _ in range(4):  # 12 MB -> two parts + remainder
            writer.write(chunk)
        writer.close()
        assert s3.get_size("s3://bkt/big.bin") == 12 * 1024 * 1024
        assert _open_uploads(s3) == []

    def test_csv_writer_aborts_the_upload_when_the_body_fails(self, s3):
        handler = UnifiedIOHandler(s3_client=s3)
        with pytest.raises(RuntimeError):
            with UnifiedCSVWriter(handler, "s3://bkt/out.csv") as writer:
                writer.writerow(["a", "b"])
                raise RuntimeError("boom")
        assert not s3.exists("s3://bkt/out.csv")


class TestS3Paths:
    @pytest.mark.parametrize(
        "uri,key",
        [
            ("s3://b/reports/q1?final.csv", "reports/q1?final.csv"),
            ("s3://b/reports/q1#final.csv", "reports/q1#final.csv"),
            ("s3://b//leading/slash.csv", "/leading/slash.csv"),
            ("s3://b/a%20b.csv", "a%20b.csv"),
            ("s3://b/", ""),
            ("s3://b", ""),
        ],
    )
    def test_key_is_everything_after_the_bucket(self, uri, key):
        path = S3Path(uri)
        assert path.bucket == "b"
        assert path.key == key

    def test_invalid_buckets(self):
        for bad in ("s3:///key", "s3://bad bucket/k", "s3://b?x/k", "s3://b#x/k"):
            with pytest.raises(ValueError):
                S3Path(bad)

    def test_path_objects_are_recognised(self):
        """Path("s3://b/k") collapses to "s3:/b/k"."""
        assert str(Path("s3://b/k")) == "s3:/b/k"
        assert is_s3_path(Path("s3://b/k")) is True
        assert is_s3_path("s3://b/k") is True
        assert is_s3_path(S3Path("s3://b/k")) is True
        assert is_s3_path("s3:/b/k") is False  # a plain string is taken literally
        assert is_s3_path(Path("local/file.csv")) is False
        assert is_s3_path(None) is False
        assert is_s3_path(42) is False

        assert normalize_s3_uri(Path("s3://b/k")) == "s3://b/k"
        assert normalize_s3_uri("s3://b/k") == "s3://b/k"
        assert S3Path(Path("s3://bkt/dir/file.csv")).key == "dir/file.csv"

    def test_client_accepts_path_objects(self, s3):
        s3._s3_client.put_object(Bucket="bkt", Key="dir/file.txt", Body=b"hi")
        assert s3.exists(Path("s3://bkt/dir/file.txt"))
        assert s3.get_size(Path("s3://bkt/dir/file.txt")) == 2
        assert UnifiedIOHandler(s3_client=s3).exists(Path("s3://bkt/dir/file.txt"))

    def test_join_and_parent_keep_working(self):
        assert str(S3Path("s3://b/base").join("x", "y?z.csv")) == "s3://b/base/x/y?z.csv"
        assert S3Path("s3://b/file.csv").parent.key == ""

    def test_list_objects_honours_max_keys_as_a_limit(self, s3):
        for index in range(5):
            s3._s3_client.put_object(Bucket="bkt", Key=f"p/{index}.txt", Body=b"x")
        assert len(list(s3.list_objects("s3://bkt/p/", max_keys=2))) == 2
        assert len(list(s3.list_objects("s3://bkt/p/", max_keys=5))) == 5
        assert len(list(s3.list_objects("s3://bkt/p/"))) == 5
        with pytest.raises(ValueError):
            list(s3.list_objects("s3://bkt/p/", max_keys=0))


class TestUnifiedIO:
    def test_open_for_read_binary_local_and_s3(self, s3, tmp_path):
        payload = b"\x00\x01PAR1\xff\r\n"
        local = tmp_path / "f.bin"
        local.write_bytes(payload)
        s3._s3_client.put_object(Bucket="bkt", Key="f.bin", Body=payload)
        handler = UnifiedIOHandler(s3_client=s3)

        for path in (local, "s3://bkt/f.bin"):
            for kwargs in ({"encoding": "binary"}, {"mode": "rb"}):
                with handler.open_for_read(path, **kwargs) as f:
                    assert f.read() == payload
                    f.seek(0)  # readers such as pyarrow/openpyxl need random access
                    assert f.read(2) == b"\x00\x01"

    def test_s3_binary_read_can_be_left_forward_only(self, s3):
        s3._s3_client.put_object(Bucket="bkt", Key="f.bin", Body=b"abc")
        handler = UnifiedIOHandler(s3_client=s3)
        with handler.open_for_read("s3://bkt/f.bin", mode="rb", seekable=False) as stream:
            assert stream.read() == b"abc"

    def test_parquet_can_be_read_from_s3_through_the_binary_mode(self, s3):
        buffer = io.BytesIO()
        pq.write_table(pa.table({"a": [1, 2, 3]}), buffer)
        s3._s3_client.put_object(Bucket="bkt", Key="t.parquet", Body=buffer.getvalue())
        handler = UnifiedIOHandler(s3_client=s3)
        with handler.open_for_read("s3://bkt/t.parquet", encoding="binary") as f:
            assert pq.ParquetFile(f).read().to_pydict() == {"a": [1, 2, 3]}

    def test_text_mode_preserves_crlf_inside_quoted_csv_fields(self, s3, tmp_path):
        data = 'id,note\r\n1,"line1\r\nline2"\r\n2,plain\r\n'
        local = tmp_path / "q.csv"
        local.write_bytes(data.encode())
        s3._s3_client.put_object(Bucket="bkt", Key="q.csv", Body=data.encode())
        handler = UnifiedIOHandler(s3_client=s3)

        for path in (local, "s3://bkt/q.csv"):
            rows = list(handler.csv_reader(path))
            assert rows == [["id", "note"], ["1", "line1\r\nline2"], ["2", "plain"]]

    def test_csv_reader_strips_a_utf8_bom(self, s3, tmp_path):
        data = b'\xef\xbb\xbf"id","name"\n1,a\n'
        local = tmp_path / "bom.csv"
        local.write_bytes(data)
        s3._s3_client.put_object(Bucket="bkt", Key="bom.csv", Body=data)
        handler = UnifiedIOHandler(s3_client=s3)

        for path in (local, "s3://bkt/bom.csv"):
            assert list(handler.csv_reader(path))[0] == ["id", "name"]

    def test_copy_file_onto_itself_does_not_truncate(self, tmp_path):
        path = tmp_path / "keep.txt"
        path.write_bytes(b"precious")
        handler = UnifiedIOHandler()
        with pytest.raises(shutil.SameFileError):
            handler.copy_file(path, path)
        with pytest.raises(shutil.SameFileError):
            handler.copy_file(path, tmp_path / "." / "keep.txt")
        assert path.read_bytes() == b"precious"

    def test_copy_file_is_binary_safe_between_all_locations(self, s3, tmp_path):
        payload = b"a\r\nb\r\n\xff\xfe\x00binary\n\r"
        src = tmp_path / "src.bin"
        src.write_bytes(payload)
        handler = UnifiedIOHandler(s3_client=s3)

        handler.copy_file(src, tmp_path / "sub" / "copy.bin")  # local -> local
        assert (tmp_path / "sub" / "copy.bin").read_bytes() == payload

        handler.copy_file(src, "s3://bkt/up.bin")  # local -> S3
        assert _object_body(s3, "up.bin") == payload

        handler.copy_file("s3://bkt/up.bin", "s3://bkt/copy.bin")  # S3 -> S3 (managed copy)
        assert _object_body(s3, "copy.bin") == payload

        handler.copy_file("s3://bkt/copy.bin", tmp_path / "down.bin")  # S3 -> local
        assert (tmp_path / "down.bin").read_bytes() == payload

        with pytest.raises(shutil.SameFileError):
            handler.copy_file("s3://bkt/up.bin", "s3://bkt/up.bin")

    def test_failed_local_copy_leaves_no_partial_destination(self, tmp_path):
        class Exploding(io.BytesIO):
            def read(self, size=-1):
                raise OSError("disk gone")

        handler = UnifiedIOHandler()
        dest = tmp_path / "dest.bin"
        dest.write_bytes(b"old")
        with patch.object(handler, "open_for_read", return_value=Exploding(b"x")):
            with pytest.raises(OSError):
                handler.copy_file(tmp_path / "src", dest)
        assert dest.read_bytes() == b"old"
        assert [p.name for p in tmp_path.iterdir()] == ["dest.bin"]

    def test_open_for_write_binary_mode(self, s3, tmp_path):
        handler = UnifiedIOHandler(s3_client=s3)
        with handler.open_for_write(tmp_path / "d" / "b.bin", mode="wb") as f:
            f.write(b"\x00\xff")
        assert (tmp_path / "d" / "b.bin").read_bytes() == b"\x00\xff"
        with handler.open_for_write("s3://bkt/b.bin", mode="wb") as f:
            f.write(b"\x00\xff")
        assert _object_body(s3, "b.bin") == b"\x00\xff"


class TestS3ParquetWriter:
    SCHEMA = pa.schema([("a", pa.int64())])

    def test_close_is_idempotent_and_cleans_the_temp_file(self, s3):
        writer = S3ParquetWriter("s3://bkt/t.parquet", self.SCHEMA, s3_client=s3)
        writer.write_table(pa.table({"a": [1, 2]}))
        temp_path = writer._temp_path
        writer.close()
        writer.close()  # used to raise FileNotFoundError
        assert not temp_path.exists()
        assert pq.read_table(io.BytesIO(_object_body(s3, "t.parquet"))).num_rows == 2

    def test_exception_in_with_body_uploads_nothing(self, s3):
        with pytest.raises(RuntimeError):
            with S3ParquetWriter("s3://bkt/t.parquet", self.SCHEMA, s3_client=s3) as writer:
                writer.write_table(pa.table({"a": [1]}))
                temp_path = writer._temp_path
                raise RuntimeError("boom")
        assert not s3.exists("s3://bkt/t.parquet")
        assert not temp_path.exists()

    def test_abort_is_idempotent(self, s3):
        writer = S3ParquetWriter("s3://bkt/t.parquet", self.SCHEMA, s3_client=s3)
        temp_path = writer._temp_path
        writer.abort()
        writer.abort()
        writer.close()
        assert not temp_path.exists()
        assert not s3.exists("s3://bkt/t.parquet")

    def test_failing_upload_still_removes_the_temp_file(self, s3):
        writer = S3ParquetWriter("s3://bkt/t.parquet", self.SCHEMA, s3_client=s3)
        temp_path = writer._temp_path
        with patch.object(s3._s3_client, "upload_fileobj", side_effect=OSError("network")):
            with pytest.raises(OSError):
                writer.close()
        assert not temp_path.exists()

    def test_constructor_failure_does_not_leak_the_temp_file(self, s3):
        created = []

        def failing_writer(path, *args, **kwargs):
            created.append(Path(path))
            raise OSError("cannot open writer")

        with patch("forklift.io.unified_io.pq.ParquetWriter", side_effect=failing_writer):
            with pytest.raises(OSError):
                S3ParquetWriter("s3://bkt/t.parquet", self.SCHEMA, s3_client=s3)
        assert created and not created[0].exists()


# =====================================================================================
# 5. Fixed-width files
# =====================================================================================


def _fwf_config(**kwargs) -> FwfInputConfig:
    fields = kwargs.pop(
        "fields",
        [
            FwfFieldSpec("id", 1, 3, parquet_type="int64"),
            FwfFieldSpec("name", 4, 6, parquet_type="string"),
        ],
    )
    return FwfInputConfig(fields=fields, **kwargs)


def _write(tmp_path: Path, text: str, name: str = "f.txt", encoding: str = "utf-8") -> Path:
    path = tmp_path / name
    path.write_bytes(text.encode(encoding))
    return path


class TestFwfRegexes:
    def test_malformed_comment_pattern_is_a_config_error(self):
        with pytest.raises(ValueError, match="comment_patterns"):
            FwfInputConfig(fields=[FwfFieldSpec("a", 1, 1)], comment_patterns=["(unclosed"])

    def test_malformed_footer_pattern_is_a_config_error(self):
        with pytest.raises(ValueError, match="footer_detection"):
            FwfInputConfig(
                fields=[FwfFieldSpec("a", 1, 1)],
                footer_detection={"mode": "regex", "pattern": "[bad"},
            )

    def test_pattern_broken_after_construction_fails_loudly_not_silently(self, tmp_path):
        """Used to yield [] with success reported: every line failed and was skipped."""
        config = _fwf_config()
        handler = FwfInputHandler(config)
        config.comment_patterns = ["(unclosed"]
        path = _write(tmp_path, "001alice\n002bob  \n")
        with pytest.raises(ValueError, match="comment_patterns"):
            handler.read_file(path)

    def test_valid_patterns_still_skip_lines(self, tmp_path):
        config = _fwf_config(
            comment_patterns=["^#"], footer_detection={"mode": "regex", "pattern": "^TOTAL"}
        )
        path = _write(tmp_path, "# note\n001alice\nTOTAL 1\n")
        records = FwfInputHandler(config).read_file(path)
        assert [r["id"] for r in records] == [1]


class TestFwfConversion:
    def test_bool_only_accepts_known_tokens(self):
        for token in ("xyz", "N/A", "maybe", "2"):
            assert FwfTypeConverter.convert_value_checked(token, "bool") == (None, False)
            assert FwfTypeConverter.convert_value(token, "bool") is None
        for token in ("TRUE", "y", " 1 ", "t", "Yes"):
            assert FwfTypeConverter.convert_value(token, "bool") is True
        for token in ("false", "N", "0", "no", "f"):
            assert FwfTypeConverter.convert_value(token, "bool") is False

    def test_bad_numbers_become_none_and_are_flagged(self):
        assert FwfTypeConverter.convert_value_checked("12A", "int64") == (None, False)
        assert FwfTypeConverter.convert_value_checked("1.5.2", "float64") == (None, False)
        assert FwfTypeConverter.convert_value_checked("300", "int8") == (None, False)
        assert FwfTypeConverter.convert_value_checked("-1", "uint8") == (None, False)
        assert FwfTypeConverter.convert_value_checked("42", "int64") == (42, True)
        assert FwfTypeConverter.convert_value_checked("", "int64") == (None, True)

    def test_errors_are_recorded_with_line_and_field_but_no_values(self, tmp_path):
        config = _fwf_config(
            fields=[
                FwfFieldSpec("id", 1, 3, parquet_type="int64"),
                FwfFieldSpec("flag", 4, 3, parquet_type="bool"),
            ]
        )
        handler = FwfInputHandler(config)
        path = _write(tmp_path, "001yes\n12Axyz\n003no \n")

        records = handler.read_file(path)

        assert [(r["id"], r["flag"]) for r in records] == [(1, True), (None, None), (3, False)]
        assert handler.errors == [
            {"line_number": 2, "field": "id", "error": "invalid_value", "type": "int64"},
            {"line_number": 2, "field": "flag", "error": "invalid_value", "type": "bool"},
        ]
        assert "12A" not in str(handler.errors) and "xyz" not in str(handler.errors)

        handler.read_file(path)  # errors describe the last read, they don't accumulate
        assert len(handler.errors) == 2

    def test_required_field_rejects_the_line(self, tmp_path):
        config = _fwf_config(
            fields=[
                FwfFieldSpec("id", 1, 3, parquet_type="int64", required=True),
                FwfFieldSpec("name", 4, 6, parquet_type="string"),
            ]
        )
        handler = FwfInputHandler(config)
        records = handler.read_file(_write(tmp_path, "001alice\n   bob   \n"))
        assert [r["id"] for r in records] == [1]
        assert handler.rejected_lines == [
            {"line_number": 2, "field": "id", "error": "required_missing"}
        ]

    def test_global_trim_whitespace_is_honoured(self):
        field = FwfFieldSpec("name", 1, 6)
        assert FwfInputHandler(_fwf_config(fields=[field])).parse_line(" ab   ") == {"name": "ab"}
        untrimmed = FwfInputHandler(_fwf_config(fields=[field], trim_whitespace=False))
        assert untrimmed.parse_line(" ab   ") == {"name": " ab   "}

    def test_field_level_trim_is_still_honoured(self):
        field = FwfFieldSpec("name", 1, 6, trim=False)
        assert FwfInputHandler(_fwf_config(fields=[field])).parse_line(" ab   ") == {
            "name": " ab   "
        }


class TestFwfBom:
    def test_bom_on_the_first_line_does_not_shift_columns(self, tmp_path):
        path = _write(tmp_path, "﻿001alice\n002bob\n")
        records = FwfInputHandler(_fwf_config()).read_file(path)
        assert [(r["id"], r["name"]) for r in records] == [(1, "alice"), (2, "bob")]

    def test_bom_with_auto_encoding(self, tmp_path):
        path = _write(tmp_path, "﻿001alice\n")
        config = _fwf_config(encoding="auto")
        assert FwfInputHandler(config).read_file(path)[0]["id"] == 1


class TestFwfValidation:
    @pytest.mark.parametrize(
        "field,message",
        [
            (FwfFieldSpec("a", 0, 3), "greater than 0"),
            (FwfFieldSpec("a", -2, 3), "negative"),
            (FwfFieldSpec("a", 1, 0), "length must be greater than 0"),
            (FwfFieldSpec("a", 1, -3), "length must be greater than 0"),
            (FwfFieldSpec("a", 1, 3, parquet_type="unknown_type"), "Invalid data type"),
            (FwfFieldSpec("", 1, 3), "cannot be empty"),
        ],
    )
    def test_simple_fields_are_validated(self, field, message):
        with pytest.raises(ValueError, match=message):
            FwfInputHandler(FwfInputConfig(fields=[field]))

    def test_duplicate_names_in_simple_fields(self):
        fields = [FwfFieldSpec("a", 1, 2), FwfFieldSpec("a", 3, 2)]
        with pytest.raises(ValueError, match="Duplicate field name"):
            FwfInputHandler(FwfInputConfig(fields=fields))

    def test_valid_aliases_are_still_accepted(self):
        fields = [
            FwfFieldSpec("a", 1, 2, parquet_type="double"),
            FwfFieldSpec("b", 3, 2, parquet_type="utf8"),
            FwfFieldSpec("c", 5, 2, parquet_type="decimal128(10,2)"),
        ]
        FwfInputHandler(FwfInputConfig(fields=fields))


class TestFwfConditional:
    @staticmethod
    def _config() -> FwfInputConfig:
        return FwfInputConfig(
            flag_column=FwfFieldSpec("rec", 1, 1),
            conditional_schemas=[
                # neither schema lists the flag column among its own fields
                FwfConditionalSchema("H", "header", [FwfFieldSpec("title", 2, 5)]),
                FwfConditionalSchema(
                    "D", "detail", [FwfFieldSpec("qty", 2, 3, parquet_type="int64")]
                ),
            ],
        )

    def test_flag_column_is_populated(self, tmp_path):
        handler = FwfInputHandler(self._config())
        records = handler.read_file(_write(tmp_path, "HHello\nD042\n"))
        assert [r["rec"] for r in records] == ["H", "D"]

        table = handler.create_arrow_table(_write(tmp_path, "HHello\nD042\n", "t.txt"))
        assert table["rec"].to_pylist() == ["H", "D"]
        assert table["qty"].to_pylist() == [None, 42]

    def test_lines_matching_no_schema_are_recorded_as_rejected(self, tmp_path):
        handler = FwfInputHandler(self._config())
        records = handler.read_file(_write(tmp_path, "D001\nXjunk\nD002\nZmore\n"))
        assert len(records) == 2
        assert handler.rejected_lines == [
            {"line_number": 2, "error": "no_matching_schema"},
            {"line_number": 4, "error": "no_matching_schema"},
        ]

    def test_parse_line_keeps_its_signature(self):
        handler = FwfInputHandler(self._config())
        assert handler.parse_line("D007") == {"qty": 7, "rec": "D"}
        assert handler.parse_line("X007") is None


# =====================================================================================
# 6./7. Encoding detection and CSV header search
# =====================================================================================


class TestEncodingDetection:
    def test_detector_reporting_no_encoding_falls_back_to_utf8(self, tmp_path):
        """chardet answers {'encoding': None, ...}: the key exists, so .get(default) was None."""
        path = _write(tmp_path, "a,b\n1,2\n")
        with patch("chardet.detect", return_value={"encoding": None, "confidence": 0.0}):
            assert CsvInputHandler(CsvInputConfig()).detect_encoding(path) == "utf-8"
            from forklift.inputs.fwf import FwfEncodingDetector

            assert FwfEncodingDetector.detect_encoding(path) == "utf-8"

    def test_non_ascii_after_an_ascii_head_is_detected_by_verifying_the_whole_file(self, tmp_path):
        head = b"id,name\n" + b"1,plain ascii row\n" * 20000  # far more than the sample
        raw = head + "2,café\n".encode("latin-1")
        path = tmp_path / "late.csv"
        path.write_bytes(raw)

        with patch("chardet.detect", return_value={"encoding": "ascii", "confidence": 1.0}):
            encoding = detect_encoding(path)
        assert encoding.lower() not in ("ascii", "utf-8")
        assert raw.decode(encoding).endswith("café\n")  # the whole file decodes

    def test_open_text_auto_actually_validates_the_encoding(self, tmp_path):
        """A latin-1 file used to come back as a utf-8-sig handle that failed on read()."""
        path = tmp_path / "latin.txt"
        path.write_bytes("naïve café\n".encode("latin-1"))
        with open_text_auto(str(path)) as handle:
            assert handle.read() == "naïve café\n"

    def test_open_text_auto_keeps_utf8_and_strips_a_bom(self, tmp_path):
        path = tmp_path / "bom.txt"
        path.write_bytes(b"\xef\xbb\xbfplain\n")
        with open_text_auto(str(path)) as handle:
            assert handle.read() == "plain\n"

    def test_open_text_auto_falls_back_to_replacement_when_nothing_fits(self, tmp_path):
        path = tmp_path / "x.txt"
        path.write_bytes(b"abc\xff")
        with open_text_auto(str(path), encodings=["utf-8", "ascii"]) as handle:
            assert handle.read() == "abc�"

    def test_works_without_any_detector_library(self, tmp_path):
        real_import = builtins.__import__

        def no_detectors(name, *args, **kwargs):
            if name in ("chardet", "charset_normalizer"):
                raise ImportError(name)
            return real_import(name, *args, **kwargs)

        utf8 = tmp_path / "u.txt"
        utf8.write_bytes("café\n".encode("utf-8"))
        latin = tmp_path / "l.txt"
        latin.write_bytes("café\n".encode("latin-1"))
        with patch("builtins.__import__", side_effect=no_detectors):
            assert detect_encoding(utf8) == "utf-8"
            chosen = detect_encoding(latin)
        assert latin.read_bytes().decode(chosen) == "café\n"

    def test_charset_normalizer_is_used_when_chardet_is_missing(self, tmp_path):
        path = _write(tmp_path, "abc\n")
        real_import = builtins.__import__

        def only_charset_normalizer(name, *args, **kwargs):
            if name == "chardet":
                raise ImportError(name)
            return real_import(name, *args, **kwargs)

        fake = types.SimpleNamespace(
            from_bytes=lambda sample: types.SimpleNamespace(
                best=lambda: types.SimpleNamespace(encoding="cp1252")
            )
        )
        with patch.dict(sys.modules, {"charset_normalizer": fake}):
            with patch("builtins.__import__", side_effect=only_charset_normalizer):
                assert detect_encoding(path) == "cp1252"

    def test_boms_decide_immediately(self, tmp_path):
        assert detect_encoding(_write(tmp_path, "x", "a.txt", "utf-8-sig")) == "utf-8-sig"
        assert detect_encoding(_write(tmp_path, "x", "b.txt", "utf-16")) == "utf-16"

    def test_verify_encoding_handles_split_multibyte_sequences(self, tmp_path):
        path = tmp_path / "multi.txt"
        path.write_bytes("€".encode("utf-8") * 10)  # 3-byte characters
        assert verify_encoding(path, "utf-8", chunk_size=4) is True
        assert verify_encoding(path, "ascii") is False
        assert verify_encoding(path, "not-a-codec") is False


class TestCsvHeaderSearch:
    def test_header_mode_absent_means_no_header(self, tmp_path):
        path = _write(tmp_path, "1,2\n3,4\n", "a.csv")
        handler = CsvInputHandler(CsvInputConfig(header_mode="absent"))
        assert handler.find_header_row(path) == (-1, [])

    def test_header_mode_auto_prefers_a_text_row(self, tmp_path):
        path = _write(tmp_path, "1,2,3\nid,name,qty\n4,5,6\n", "a.csv")
        handler = CsvInputHandler(CsvInputConfig(header_mode="auto"))
        assert handler.find_header_row(path) == (1, ["id", "name", "qty"])
        # present mode takes the first non-blank row as is
        assert CsvInputHandler(CsvInputConfig()).find_header_row(path) == (0, ["1", "2", "3"])

    def test_quote_char_is_honoured(self, tmp_path):
        path = _write(tmp_path, "'a,b',c\n1,2\n", "a.csv")
        handler = CsvInputHandler(CsvInputConfig(quote_char="'"))
        assert handler.find_header_row(path) == (0, ["a,b", "c"])

    def test_escape_char_is_honoured(self, tmp_path):
        path = _write(tmp_path, "a\\,b,c\n1,2\n", "a.csv")
        handler = CsvInputHandler(CsvInputConfig(escape_char="\\"))
        assert handler.find_header_row(path) == (0, ["a,b", "c"])

    def test_newlines_inside_quoted_header_cells_survive(self, tmp_path):
        path = _write(tmp_path, '"na\r\nme",x\r\n1,2\r\n', "a.csv")
        _, names = CsvInputHandler(CsvInputConfig()).find_header_row(path)
        assert names == ["na\r\nme", "x"]

    def test_bom_is_not_part_of_the_first_header_name(self, tmp_path):
        path = _write(tmp_path, "﻿id,name\n1,a\n", "a.csv")
        assert CsvInputHandler(CsvInputConfig()).find_header_row(path)[1] == ["id", "name"]

    def test_blank_first_line_is_never_the_header(self, tmp_path):
        path = _write(tmp_path, "\nid,name\n", "a.csv")
        config = CsvInputConfig(skip_blank_lines=False)
        assert CsvInputHandler(config).find_header_row(path) == (1, ["id", "name"])

    def test_only_blank_lines_raise(self, tmp_path):
        path = _write(tmp_path, "\n\n", "a.csv")
        with pytest.raises(ValueError, match="No valid header row"):
            CsvInputHandler(CsvInputConfig()).find_header_row(path)

    def test_comment_patterns_are_compiled_up_front(self):
        with pytest.raises(ValueError, match="comment_patterns"):
            CsvInputConfig(comment_patterns=["[bad"])
        config = CsvInputConfig()
        config.comment_patterns = ["[bad"]
        with pytest.raises(ValueError, match="comment_patterns"):
            CsvInputHandler(config)

    def test_comment_rows_are_still_skipped(self, tmp_path):
        path = _write(tmp_path, "# note\nid,name\n", "a.csv")
        handler = CsvInputHandler(CsvInputConfig(comment_patterns=[r"^#"]))
        assert handler.find_header_row(path) == (1, ["id", "name"])
