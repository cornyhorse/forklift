"""run_job: every kind of job, every location, and how failures become results.

A job never raises for its own failures; these tests check the result a caller gets in each
case: the status, the error code and message (no cell values, no secrets), the counts, the
artifacts (with rows, bytes and sha256) and the warnings.
"""

from __future__ import annotations

import hashlib
import io
import json
import os
import sys
import types
from pathlib import Path
from unittest.mock import MagicMock, patch

import jsonschema
import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from forklift.engine.exceptions import ImportCancelled, LimitExceededError
from forklift.jobs import JobSpec, contract, run_job
from forklift.jobs import runner as runner_module

RESULT_VALIDATOR = jsonschema.Draft202012Validator(contract.jobresult_schema())
PEOPLE = "id,name,age\n1,Ana,34\n2,Bo,x\n3,Cy,29\n"
SCHEMA = {
    "properties": {
        "id": {"type": "integer"},
        "name": {"type": "string"},
        "age": {"type": "integer"},
    },
    "required": ["id"],
}


def _spec(kind="run", fmt="csv", location=None, **extra):
    spec = {
        "spec_version": 1,
        "job_id": "job-1",
        "kind": kind,
        "input": {
            "format": fmt,
            "location": location or {"type": "file", "path": "in/people.csv"},
            "options": extra.pop("input_options", {}),
        },
        "output": {"location": {"type": "file", "path": "out/"}},
    }
    spec.update(extra)
    return spec


@pytest.fixture
def base(tmp_path):
    (tmp_path / "in").mkdir()
    (tmp_path / "in" / "people.csv").write_text(PEOPLE)
    return tmp_path


def _run(spec, base, **kwargs):
    result = run_job(spec, base_dir=base, **kwargs)
    assert RESULT_VALIDATOR.is_valid(result.to_dict()), list(
        RESULT_VALIDATOR.iter_errors(result.to_dict())
    )
    return result


def _sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _workbook(path, sheets=None):
    workbook = openpyxl.Workbook()
    sheets = sheets or {"people": [["id", "name"], [1, "Ana"], [2.0, "Bo"]]}
    for index, (name, rows) in enumerate(sheets.items()):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name
        for row in rows:
            sheet.append(row)
    workbook.save(path)
    return path


class TestArguments:
    def test_base_dir_must_be_a_directory(self, tmp_path):
        with pytest.raises(ValueError, match="base_dir is not a directory"):
            run_job(_spec(), base_dir=tmp_path / "missing")

    def test_allowed_hosts_must_not_be_a_string(self, base):
        with pytest.raises(ValueError, match="list of host names"):
            run_job(_spec(), base_dir=base, allowed_url_hosts="store")

    def test_callbacks_must_be_callable(self, base):
        with pytest.raises(TypeError, match="progress must be callable"):
            run_job(_spec(), base_dir=base, progress=1)

    @pytest.mark.parametrize(
        "document, job_id",
        [({"job_id": "j-9", "kind": "nope"}, "j-9"), ({"job_id": ""}, None), ([], None)],
    )
    def test_invalid_spec_is_a_failed_result(self, base, document, job_id):
        result = _run(document, base)
        assert result.status == "failed" and result.job_id == job_id
        assert result.error.code == "SPEC_INVALID" and not result.error.retryable
        assert result.error.message.startswith("Invalid job spec")

    def test_a_jobspec_object_runs_too(self, base):
        result = _run(JobSpec.from_dict(_spec()), base)
        assert result.status == "succeeded"


class TestRunCsv:
    def test_run_writes_parquet_and_describes_every_file(self, base):
        events = []
        result = _run(_spec(schema=SCHEMA), base, progress=events.append)

        assert result.status == "succeeded" and result.error is None
        assert result.counts == {
            "total_rows": 3,
            "valid_rows": 2,
            "invalid_rows": 1,
            "truncated_rows": 0,
        }
        kinds = [(a.kind, a.path) for a in result.artifacts]
        assert kinds == [
            ("data", "out/data.parquet"),
            ("bad_rows", "out/bad_rows.parquet"),
            ("metadata", "out/output_data_metadata.json"),
            ("metadata", "out/metadata.json"),
            ("manifest", "out/manifest.json"),
        ]
        data, bad = result.artifacts[0], result.artifacts[1]
        assert (data.rows, bad.rows) == (2, 1)
        assert data.bytes == (base / "out" / "data.parquet").stat().st_size
        assert data.sha256 == _sha256(base / "out" / "data.parquet")
        assert result.artifacts[3].rows is None
        assert events == [{"rows_read": 3, "rows_rejected": 1, "bytes_read": len(PEOPLE)}]
        # The private directory with the inline schema is gone
        assert sorted(p.name for p in base.iterdir()) == ["in", "out"]

    def test_options_reach_the_engine(self, base):
        (base / "in" / "semi.csv").write_text("# comment\nid;name\n1;a\n2;b\nTOTAL;2\n")
        spec = _spec(
            location={"type": "file", "path": "in/semi.csv"},
            input_options={
                "delimiter": ";",
                "footer_detection": {"column_index": 0, "patterns": ["^TOTAL$"]},
            },
            output={"location": {"type": "file", "path": "out/"}, "compression": "gzip"},
            options={"batch_size": 1, "apply_schema_extensions": False},
        )
        result = _run(spec, base)
        assert result.counts["total_rows"] == 2
        metadata = pq.read_metadata(base / "out" / "data.parquet")
        assert metadata.row_group(0).column(0).compression == "GZIP"

    def test_schema_extensions_findings_and_warnings(self, base):
        schema = dict(
            SCHEMA,
            **{
                "x-pii": {"name": {}},
                "x-uniqueConstraints": [{"name": "u_name", "columns": ["name"]}],
            },
        )
        (base / "in" / "dupes.csv").write_text("id,name,age\n1,a,1\n2,a,2\n")
        spec = _spec(location={"type": "file", "path": "in/dupes.csv"}, schema=schema)
        result = _run(spec, base)
        assert result.status == "succeeded"
        assert result.schema_extensions
        assert result.validation_summary
        assert result.warnings

    def test_threshold_failure_keeps_the_bad_rows(self, base):
        schema = dict(
            SCHEMA,
            **{
                "x-validation": {
                    "fieldValidations": {"age": {"range": {"min": 0, "max": 40}}},
                    "badRowsHandling": {"maxBadRowsPercent": 10},
                }
            },
        )
        rows = "".join(f"{i},n,{90 if i % 2 else 30}\n" for i in range(20))
        (base / "in" / "ages.csv").write_text("id,name,age\n" + rows)
        spec = _spec(location={"type": "file", "path": "in/ages.csv"}, schema=schema)

        result = _run(spec, base)

        assert result.status == "failed"
        assert result.error.code == "BAD_ROWS_THRESHOLD_EXCEEDED"
        assert [(a.kind, a.path) for a in result.artifacts] == [
            ("bad_rows", "out/bad_rows.parquet")
        ]
        assert result.artifacts[0].rows > 0 and result.counts["invalid_rows"] > 0
        assert str(base) not in result.error.message
        assert not (base / "out" / "data.parquet").exists()

    def test_engine_errors_keep_their_code_and_message(self, base):
        spec = _spec(schema={"properties": {"email": {}}, "required": ["email"]})
        result = _run(spec, base)
        assert result.error.code == "COLUMN_MISSING"
        assert "Required column(s) missing from the input header: 'email'" in result.error.message

    def test_missing_input_file(self, base):
        result = _run(_spec(location={"type": "file", "path": "in/none.csv"}), base)
        assert result.error.code == "INPUT_UNREADABLE"
        assert result.error.message == (
            "FileNotFoundError: Input file not found: in/none.csv (relative to the base directory)"
        )

    def test_symlink_out_of_the_base_directory_is_refused(self, base, tmp_path_factory):
        outside = tmp_path_factory.mktemp("outside")
        (outside / "secret.csv").write_text("a\n1\n")
        os.symlink(outside, base / "in" / "link")
        result = _run(_spec(location={"type": "file", "path": "in/link/secret.csv"}), base)
        assert result.error.code == "SPEC_INVALID"
        assert "leads outside the base directory" in result.error.message

    def test_input_larger_than_the_limit_is_not_read(self, base):
        result = _run(_spec(limits={"max_input_bytes": 10}), base)
        assert result.error.code == "LIMIT_EXCEEDED"
        assert f"The input is {len(PEOPLE)} bytes" in result.error.message
        assert not (base / "out" / "data.parquet").exists()

    def test_row_limit(self, base):
        result = _run(_spec(limits={"max_rows": 2}), base)
        assert result.error.code == "LIMIT_EXCEEDED"
        assert "more than 2 rows" in result.error.message
        assert result.counts == {"total_rows": 3, "invalid_rows": 0}

    def test_time_limit(self, base, monkeypatch):
        ticks = iter([0.0, 100.0, 100.0, 100.0])
        monkeypatch.setattr(runner_module, "clock", lambda: next(ticks))
        result = _run(_spec(limits={"max_seconds": 5}), base)
        assert result.error.code == "LIMIT_EXCEEDED"
        assert "longer than 5 seconds" in result.error.message

    def test_cancel(self, base):
        result = _run(_spec(), base, cancel=lambda: True)
        assert result.status == "cancelled"
        assert result.error.code == "CANCELLED" and not result.error.retryable
        assert result.artifacts == []

    def test_unknown_encoding_is_spec_invalid(self, base):
        result = _run(_spec(input_options={"encoding": "utf-9"}), base)
        assert result.error.code == "SPEC_INVALID"
        assert "input.options.encoding: unknown encoding 'utf-9'" in result.error.message

    def test_undecodable_input_is_an_encoding_error(self, base):
        (base / "in" / "latin.csv").write_bytes("name\nJosé\n".encode("latin-1"))
        result = _run(_spec(location={"type": "file", "path": "in/latin.csv"}), base)
        assert result.error.code == "ENCODING_ERROR"
        assert "not valid for encoding 'utf-8'" in result.error.message

    def test_options_that_do_not_apply_are_ignored_with_a_warning(self, base):
        spec = _spec(
            input_options={"sheet": "x", "query_timeout": 3},
            options={"preview_rows": 5, "infer_primary_key": True},
        )
        result = _run(spec, base)
        assert result.status == "succeeded"
        assert result.warnings == [
            "input.options.sheet only applies to format 'excel' and is ignored",
            "input.options.query_timeout only applies to format 'sql' and is ignored",
            "options.preview_rows only applies to kind 'preview' and is ignored",
            "options.infer_primary_key only applies to kind 'generate_schema' and is ignored",
        ]

    def test_fwf_is_not_implemented(self, base):
        result = _run(_spec(fmt="fwf"), base)
        assert result.error.code == "SPEC_INVALID"
        assert "fixed-width import is not implemented" in result.error.message

    def test_unexpected_errors_are_internal(self, base):
        with patch(
            "forklift.engine.forklift_core.ForkliftCore.process_csv",
            side_effect=ZeroDivisionError("division by zero"),
        ):
            result = _run(_spec(), base)
        assert result.error.code == "INTERNAL"
        assert result.error.message == "ZeroDivisionError: division by zero"


class TestRunExcel:
    def test_every_sheet_becomes_a_data_artifact(self, base):
        _workbook(base / "in" / "book.xlsx", {"a": [["id"], [1]], "b": [["id"], [2], [3]]})
        spec = _spec(
            fmt="excel",
            location={"type": "file", "path": "in/book.xlsx"},
            output={"location": {"type": "file", "path": "out/"}, "compression": "zstd"},
        )
        result = _run(spec, base)
        assert result.status == "succeeded"
        assert [(a.path, a.rows) for a in result.artifacts] == [
            ("out/book_a.parquet", 1),
            ("out/book_b.parquet", 2),
        ]
        assert result.warnings == [
            "output.compression only applies to CSV runs and is ignored (snappy is used)"
        ]

    def test_sheet_option(self, base):
        _workbook(base / "in" / "book.xlsx", {"a": [["id"], [1]], "b": [["id"], [2]]})
        spec = _spec(
            fmt="excel",
            location={"type": "file", "path": "in/book.xlsx"},
            input_options={"sheet": "b"},
        )
        result = _run(spec, base, s3_client=MagicMock())
        assert [a.path for a in result.artifacts] == ["out/book_b.parquet"]


SQL_SCHEMA = pa.schema([("id", pa.int64())])


def _fake_sql(tables=(("public", "users", None),), rows=3):
    schema_importer = MagicMock()
    schema_importer.get_table_list.return_value = list(tables)
    schema_importer.get_selected_columns.return_value = None
    handler = MagicMock()
    handler.__enter__.return_value = handler
    handler.__exit__.return_value = None
    handler.get_table_schema.return_value = SQL_SCHEMA

    def read(schema, table, columns=None):
        yield pa.record_batch([pa.array(list(range(rows)))], schema=SQL_SCHEMA)

    handler.read_table_data.side_effect = read
    return patch(
        "forklift.schema.sql_schema_importer.SqlSchemaImporter", return_value=schema_importer
    ), patch("forklift.inputs.sql.SqlInputHandler", return_value=handler)


def _sql_spec(**extra):
    return _spec(
        fmt="sql",
        location={"type": "sql", "connection_string": "Driver=x;Uid=u;Pwd=hunter22"},
        schema={"x-sql": {"tables": [{"select": {"schema": "public", "name": "users"}}]}},
        **extra,
    )


class TestRunSql:
    def test_tables_become_data_artifacts_with_the_metadata(self, base):
        schema_patch, handler_patch = _fake_sql()
        with schema_patch, handler_patch:
            result = _run(_sql_spec(options={"batch_size": 7}), base)
        assert result.status == "succeeded", result.error
        assert [(a.kind, a.path, a.rows) for a in result.artifacts] == [
            ("data", "out/public_users.parquet", 3),
            ("metadata", "out/metadata.json", None),
        ]
        assert "hunter22" not in (base / "out" / "metadata.json").read_text()

    def test_options_and_s3_client_reach_the_importer(self, base):
        with patch("forklift.engine.forklift_core.import_sql") as import_sql:
            import_sql.return_value = MagicMock(
                total_rows=0,
                valid_rows=0,
                invalid_rows=0,
                truncated_rows=0,
                schema_extensions=[],
                validation_summary={},
                warnings=[],
                errors=[],
                output_files=[],
                metadata_file=None,
                manifest_file=None,
            )
            client = MagicMock()
            _run(_sql_spec(input_options={"query_timeout": 9}), base, s3_client=client)
        kwargs = import_sql.call_args.kwargs
        assert kwargs["query_timeout"] == 9 and kwargs["s3_client"] is client
        assert kwargs["batch_size"] == 10000 and callable(kwargs["progress"])

    def test_database_errors_never_show_the_driver_text_or_password(self, base):
        class Error(Exception):
            pass

        Error.__module__ = "pyodbc"
        failure = Error(
            "28000", "[unixODBC] login failed for u with Pwd=hunter22 (18456) (SQLDriverConnect)"
        )
        schema_patch, handler_patch = _fake_sql()
        with schema_patch, handler_patch as handler:
            handler.return_value.__enter__.side_effect = failure
            result = _run(_sql_spec(), base)
        assert result.error.code == "PERMISSION_DENIED"
        assert (
            "hunter22" not in result.error.message and "login failed" not in result.error.message
        )
        assert (
            result.error.message
            == "Error (SQLSTATE 28000, invalid authorization, driver error 18456)"
        )


class _WriteResult:
    def __init__(self, rows_written, warnings=()):
        self.rows_written = rows_written
        self.warnings = list(warnings)


@pytest.fixture
def fake_writer(monkeypatch):
    """A stand-in for forklift.outputs.sql (agent T's write_table)."""
    calls = []
    behaviour = {"result": _WriteResult(2, ["mapped double to DOUBLE PRECISION"])}

    def write_table(source, connection_string, table, **kwargs):
        calls.append(dict(kwargs, source=source, connection_string=connection_string, table=table))
        kwargs["progress"]({"rows_written": 1})
        kwargs["progress"]("not a dict")
        if "error" in behaviour:
            raise behaviour["error"]
        if kwargs["cancel"]():
            raise RuntimeError("cancelled by the caller")
        return behaviour["result"]

    module = types.ModuleType("forklift.outputs.sql")
    module.write_table = write_table
    monkeypatch.setitem(sys.modules, "forklift.outputs.sql", module)
    return calls, behaviour


def _table_spec(**location):
    target = {
        "type": "sql_table",
        "connection_string": "Driver=pg;Pwd=s3cret!",
        "table": "people",
        "schema_name": "staging",
        "mode": "upsert",
        "key_columns": ["id"],
    }
    target.update(location)
    return _spec(
        schema=SCHEMA,
        output={"location": target, "artifacts": {"type": "file", "path": "parquet/"}},
    )


class TestSqlTableOutput:
    def test_validated_rows_are_loaded_and_the_parquet_stays(self, base, fake_writer):
        calls, _ = fake_writer
        events = []
        result = _run(_table_spec(), base, progress=events.append)

        assert result.status == "succeeded", result.error
        assert result.counts["rows_written"] == 2 and result.counts["valid_rows"] == 2
        assert result.warnings == ["mapped double to DOUBLE PRECISION"]
        assert [a.path for a in result.artifacts][:2] == [
            "parquet/data.parquet",
            "parquet/bad_rows.parquet",
        ]
        (call,) = calls
        assert call["source"] == str(base / "parquet" / "data.parquet")
        assert call["table"] == "people" and call["schema_name"] == "staging"
        assert call["mode"] == "upsert" and call["key_columns"] == ["id"]
        assert call["job_id"] == "job-1" and call["staging"] == "table"
        assert events[-2:] == [
            {"rows_read": 3, "rows_rejected": 1, "bytes_read": len(PEOPLE), "rows_written": 1},
            {"rows_read": 3, "rows_rejected": 1, "bytes_read": len(PEOPLE)},
        ]

    def test_artifacts_default_to_out(self, base, fake_writer):
        spec = _table_spec()
        del spec["output"]["artifacts"]
        result = _run(spec, base)
        assert result.artifacts[0].path == "out/data.parquet"

    def test_write_failure_keeps_the_artifacts_and_hides_the_password(self, base, fake_writer):
        _, behaviour = fake_writer
        behaviour["error"] = RuntimeError("could not write with Driver=pg;Pwd=s3cret!")
        result = _run(_table_spec(), base)
        assert result.error.code == "TARGET_WRITE_FAILED"
        assert "s3cret" not in result.error.message
        assert [a.kind for a in result.artifacts][0] == "data"

    def test_code_of_the_writer_is_kept(self, base, fake_writer):
        _, behaviour = fake_writer
        error = RuntimeError("no INSERT privilege on staging.people")
        error.error_code, error.retryable = "PERMISSION_DENIED", False
        behaviour["error"] = error
        assert _run(_table_spec(), base).error.code == "PERMISSION_DENIED"

    def test_cancel_during_the_write(self, base, fake_writer):
        answers = iter([False, True])  # the import batch, then the table write
        result = _run(_table_spec(), base, cancel=lambda: next(answers))
        assert result.status == "cancelled" and result.error.code == "CANCELLED"

    def test_time_limit_hidden_by_the_writer_is_still_reported(
        self, base, fake_writer, monkeypatch
    ):
        # Start, the import's batch, then the table write: past the limit
        ticks = iter([0.0, 1.0, 50.0])
        monkeypatch.setattr(runner_module, "clock", lambda: next(ticks))

        def wrapping(source, connection_string, table, **kwargs):
            try:
                kwargs["progress"]({"rows_written": 1})
            except LimitExceededError:
                raise RuntimeError("the writer wrapped it")

        sys.modules["forklift.outputs.sql"].write_table = wrapping
        result = _run(dict(_table_spec(), limits={"max_seconds": 10}), base)
        assert result.error.code == "LIMIT_EXCEEDED"
        assert "longer than 10 seconds" in result.error.message

    def test_interruptions_from_the_writer_pass_through(self, base, fake_writer):
        _, behaviour = fake_writer
        behaviour["error"] = ImportCancelled("stop")
        assert _run(_table_spec(), base).error.code == "CANCELLED"

    def test_several_data_files_cannot_go_to_one_table(self, base, fake_writer):
        _workbook(base / "in" / "book.xlsx", {"a": [["id"], [1]], "b": [["id"], [2]]})
        spec = _table_spec()
        spec["input"] = {"format": "excel", "location": {"type": "file", "path": "in/book.xlsx"}}
        spec["schema"] = None
        result = _run(spec, base)
        assert result.error.code == "TARGET_WRITE_FAILED"
        assert "loads one data file, but the import produced 2" in result.error.message

    def test_no_data_file_loads_nothing(self, base, fake_writer):
        calls, _ = fake_writer
        schema_patch, handler_patch = _fake_sql(rows=0)
        spec = _table_spec()
        spec["input"] = _sql_spec()["input"]
        spec["schema"] = _sql_spec()["schema"]
        with schema_patch, handler_patch:
            result = _run(spec, base)
        assert result.status == "succeeded" and calls == []
        assert result.counts["rows_written"] == 0
        assert "Nothing was loaded into 'people'" in result.warnings[-1]

    def test_missing_writer_module(self, base, monkeypatch):
        monkeypatch.setitem(sys.modules, "forklift.outputs.sql", None)
        result = _run(_table_spec(), base)
        assert result.error.code == "TARGET_WRITE_FAILED"
        assert "sql_table outputs need forklift.outputs.sql" in result.error.message


class TestPreview:
    def test_first_rows_as_text(self, base):
        result = _run(_spec(kind="preview", output=None, options={"preview_rows": 2}), base)

        assert result.status == "succeeded"
        assert result.counts == {"total_rows": 2}
        (artifact,) = result.artifacts
        assert (artifact.kind, artifact.path, artifact.rows) == ("preview", "out/preview.json", 2)
        preview = json.loads((base / "out" / "preview.json").read_text())
        assert preview == {
            "columns": ["id", "name", "age"],
            "rows": [["1", "Ana", "34"], ["2", "Bo", "x"]],
            "row_count": 2,
            "truncated": True,
            "limit": "rows",
            "truncated_cells": 0,
        }

    def test_header_comments_blank_rows_and_footer(self, base):
        text = "# exported\nid,name\n\n1,a\n2,b\nTOTAL,2\n3,c\n"
        (base / "in" / "f.csv").write_text(text)
        spec = _spec(
            kind="preview",
            location={"type": "file", "path": "in/f.csv"},
            input_options={"footer_detection": {"column_index": 0, "patterns": ["^TOTAL"]}},
        )
        _run(spec, base)
        preview = json.loads((base / "out" / "preview.json").read_text())
        assert preview["rows"] == [["1", "a"], ["2", "b"]] and not preview["truncated"]

    def test_absent_header_takes_the_schema_names(self, base):
        (base / "in" / "raw.csv").write_text("1,a\n2,b\n")
        spec = _spec(
            kind="preview",
            location={"type": "file", "path": "in/raw.csv"},
            input_options={"header_mode": "absent"},
            schema={"properties": {"id": {}, "name": {}}},
        )
        _run(spec, base)
        preview = json.loads((base / "out" / "preview.json").read_text())
        assert preview["columns"] == ["id", "name"] and len(preview["rows"]) == 2

    def test_byte_limit_and_long_cells(self, base):
        long = "x" * 5000
        (base / "in" / "wide.csv").write_text("a\n" + f"{long}\n" * 400)
        spec = _spec(
            kind="preview",
            location={"type": "file", "path": "in/wide.csv"},
            options={"preview_max_bytes": 20_000, "preview_rows": 1000},
        )
        _run(spec, base)
        text = (base / "out" / "preview.json").read_text()
        preview = json.loads(text)
        assert len(text.encode()) <= 20_000
        assert preview["limit"] == "bytes" and preview["truncated"]
        assert preview["truncated_cells"] == len(preview["rows"]) > 0
        assert all(len(row[0]) == 2000 for row in preview["rows"])

    def test_undecodable_rows(self, base):
        (base / "in" / "bad.csv").write_bytes(b"a\n" + b"ok\n" * 5000 + b"caf\xe9\n")
        spec = _spec(
            kind="preview",
            location={"type": "file", "path": "in/bad.csv"},
            options={"preview_rows": 10_000},
        )
        assert _run(spec, base).error.code == "ENCODING_ERROR"

    def test_undecodable_header(self, base):
        (base / "in" / "bad.csv").write_bytes(b"caf\xe9\n1\n")
        spec = _spec(kind="preview", location={"type": "file", "path": "in/bad.csv"})
        assert _run(spec, base).error.code == "ENCODING_ERROR"

    def test_excel_sheet_by_name_and_index(self, base):
        import datetime

        _workbook(
            base / "in" / "book.xlsx",
            {
                "first": [["x"], [1]],
                "people": [
                    ["id", "name", "when", "ok", "ratio"],
                    [1, "Ana", datetime.date(2024, 1, 2), True, 0.5],
                    [2, None, datetime.datetime(2024, 1, 2, 3, 4), False, 1.0],
                ],
            },
        )
        for sheet in ("people", 1, "1"):
            spec = _spec(
                kind="preview",
                fmt="excel",
                location={"type": "file", "path": "in/book.xlsx"},
                input_options={"sheet": sheet, "values_only": True},
                options={"preview_rows": 1},
            )
            assert _run(spec, base).status == "succeeded"
            preview = json.loads((base / "out" / "preview.json").read_text())
            assert preview["sheet"] == "people"
            assert preview["columns"] == ["id", "name", "when", "ok", "ratio"]
            assert preview["rows"] == [["1", "Ana", "2024-01-02T00:00:00", "TRUE", "0.5"]]
            assert preview["limit"] == "rows"

        spec = _spec(
            kind="preview", fmt="excel", location={"type": "file", "path": "in/book.xlsx"}
        )
        _run(spec, base)
        preview = json.loads((base / "out" / "preview.json").read_text())
        assert preview["sheet"] == "first" and preview["rows"] == [["1"]]

    def test_sql_and_fwf_are_not_previewed(self, base):
        assert "not sql" in _run(_sql_spec(kind="preview"), base).error.message
        assert _run(_spec(kind="preview", fmt="fwf"), base).error.code == "SPEC_INVALID"


class TestValidateSchema:
    def _spec(self, schema, **extra):
        return _spec(kind="validate_schema", schema=schema, output=None, **extra)

    def test_report_of_a_good_schema(self, base):
        schema = {"properties": {"id": {"type": "integer"}, "extra": {}}, "required": ["id"]}
        result = _run(self._spec(schema), base)

        assert result.status == "succeeded"
        assert result.counts["total_rows"] == 3
        (artifact,) = result.artifacts
        assert (artifact.kind, artifact.path) == ("report", "out/report.json")
        report = json.loads((base / "out" / "report.json").read_text())
        assert report["valid"] is True and report["error"] is None
        assert report["columns"] == ["id", "name", "age"]
        assert report["columns_not_in_schema"] == ["name", "age"]
        assert report["schema_columns_not_in_input"] == ["extra"]
        assert report["required_columns"] == ["id"]
        assert report["sample_rows"] == 3

    def test_report_of_a_bad_schema_and_a_failed_result(self, base):
        schema = {"properties": {"email": {}}, "required": ["email"]}
        result = _run(self._spec(schema), base)

        assert result.status == "failed" and result.error.code == "COLUMN_MISSING"
        assert [a.kind for a in result.artifacts] == ["report"]
        report = json.loads((base / "out" / "report.json").read_text())
        assert report["valid"] is False
        assert report["error"]["code"] == "COLUMN_MISSING"
        assert report["schema_columns_not_in_input"] == ["email"]

    def test_only_a_sample_is_checked(self, base):
        rows = "".join(f"{i},n,{i}\n" for i in range(50))
        (base / "in" / "many.csv").write_text("# note\nid,name,age\n" + rows)
        spec = self._spec(
            SCHEMA, location={"type": "file", "path": "in/many.csv"}, options={"sample_rows": 5}
        )
        result = _run(spec, base)
        assert result.counts["total_rows"] == 5

    def test_unreadable_input_still_gives_a_report(self, base):
        spec = self._spec(SCHEMA, location={"type": "file", "path": "in/none.csv"})
        result = _run(spec, base)
        assert result.error.code == "INPUT_UNREADABLE"
        report = json.loads((base / "out" / "report.json").read_text())
        assert report["columns"] is None and report["schema_columns_not_in_input"] is None

    def test_schema_without_properties(self, base):
        report_result = _run(self._spec({"title": "x"}), base)
        assert report_result.status == "succeeded"
        report = json.loads((base / "out" / "report.json").read_text())
        assert report["schema_columns"] == [] and report["columns_not_in_schema"] == [
            "id",
            "name",
            "age",
        ]

    def test_sample_copy_respects_the_byte_limit(self, base):
        (base / "in" / "big.csv").write_text("a\n" + "1\n" * 100_000)
        spec = self._spec(
            {"properties": {"a": {}}},
            location={"type": "file", "path": "in/big.csv"},
            options={"sample_rows": 100_000},
            limits={"max_input_bytes": 1000},
        )
        assert _run(spec, base).error.code == "LIMIT_EXCEEDED"

    def test_undecodable_sample(self, base):
        (base / "in" / "bad.csv").write_bytes(b"a\n" + b"1\n" * 5000 + b"\xe9\n")
        spec = self._spec(
            {"properties": {"a": {}}},
            location={"type": "file", "path": "in/bad.csv"},
            options={"sample_rows": 10_000},
        )
        assert _run(spec, base).error.code == "ENCODING_ERROR"

    def test_other_formats_are_refused(self, base):
        result = _run(_spec(kind="validate_schema", fmt="excel", schema={}, output=None), base)
        assert result.error.code == "SPEC_INVALID"
        assert "checks a schema against a CSV input" in result.error.message


class TestGenerateSchema:
    def test_csv_schema(self, base):
        spec = _spec(
            kind="generate_schema",
            output=None,
            input_options={"delimiter": ",", "header_mode": "auto"},
            options={"sample_rows": 2, "infer_primary_key": True},
        )
        result = _run(spec, base)

        assert result.status == "succeeded"
        assert result.counts == {"total_rows": 2}
        assert result.warnings == [
            "input.options.header_mode is not used by generate_schema (the schema generator "
            "reads the first row as the header) and is ignored"
        ]
        schema = json.loads((base / "out" / "schema.json").read_text())
        assert list(schema["properties"]) == ["id", "name", "age"]
        assert schema["x-generation"]["source_file"] == "people.csv"

    def test_excel_schema(self, base):
        _workbook(base / "in" / "book.xlsx")
        spec = _spec(
            kind="generate_schema",
            fmt="excel",
            location={"type": "file", "path": "in/book.xlsx"},
            input_options={"sheet": "people"},
        )
        result = _run(spec, base)
        assert result.status == "succeeded", result.error
        assert "x-excel" in json.loads((base / "out" / "schema.json").read_text())

    def test_unreadable_inputs(self, base):
        (base / "in" / "bad.csv").write_bytes(b"a\n\xe9\n")
        (base / "in" / "empty.csv").write_bytes(b"")
        for name, code in (("bad.csv", "ENCODING_ERROR"), ("empty.csv", "INPUT_UNREADABLE")):
            spec = _spec(kind="generate_schema", location={"type": "file", "path": f"in/{name}"})
            assert _run(spec, base).error.code == code
        with patch(
            "forklift.schema.schema_generator.SchemaGenerator.generate_schema",
            side_effect=UnicodeDecodeError("utf-8", b"\xe9", 0, 1, "invalid"),
        ):
            assert _run(_spec(kind="generate_schema"), base).error.code == "ENCODING_ERROR"

    def test_without_a_row_count(self, base):
        with patch(
            "forklift.schema.schema_generator.SchemaGenerator.generate_schema",
            return_value={"properties": {}},
        ):
            result = _run(_spec(kind="generate_schema"), base)
        assert result.status == "succeeded" and result.counts == {}

    def test_sql_and_fwf(self, base):
        assert "not sql" in _run(_sql_spec(kind="generate_schema"), base).error.message
        assert _run(_spec(kind="generate_schema", fmt="fwf"), base).error.code == "SPEC_INVALID"


moto = pytest.importorskip("moto")


@pytest.fixture
def s3(monkeypatch):
    import boto3

    from forklift.io import S3StreamingClient

    monkeypatch.setenv("AWS_ACCESS_KEY_ID", "test")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "test")
    monkeypatch.setenv("AWS_DEFAULT_REGION", "us-east-1")
    with moto.mock_aws():
        raw = boto3.client("s3", region_name="us-east-1")
        raw.create_bucket(Bucket="bkt")
        raw.put_object(Bucket="bkt", Key="in/people.csv", Body=PEOPLE.encode())
        buffer = io.BytesIO()
        workbook = openpyxl.Workbook()
        workbook.active.append(["id"])
        workbook.active.append([1])
        workbook.save(buffer)
        raw.put_object(Bucket="bkt", Key="in/book.xlsx", Body=buffer.getvalue())
        yield raw, S3StreamingClient()


S3_IN = {"type": "s3", "uri": "s3://bkt/in/people.csv"}


class TestS3Locations:
    def test_csv_from_and_to_s3(self, base, s3):
        raw, client = s3
        spec = _spec(
            location=S3_IN,
            schema=SCHEMA,
            output={"location": {"type": "s3", "uri": "s3://bkt/out/"}},
            limits={"max_input_bytes": 10_000},
        )
        result = _run(spec, base, s3_client=client)
        assert result.status == "succeeded", result.error
        assert [(a.kind, a.path, a.rows, a.sha256) for a in result.artifacts[:2]] == [
            ("data", "s3://bkt/out/data.parquet", 2, None),
            ("bad_rows", "s3://bkt/out/bad_rows.parquet", 1, None),
        ]
        keys = {o["Key"] for o in raw.list_objects_v2(Bucket="bkt", Prefix="out/")["Contents"]}
        assert {"out/data.parquet", "out/manifest.json"} <= keys

    def test_s3_input_larger_than_the_limit(self, base, s3):
        result = _run(_spec(location=S3_IN, limits={"max_input_bytes": 5}), base, s3_client=s3[1])
        assert result.error.code == "LIMIT_EXCEEDED"

    def test_excel_to_s3_has_no_per_file_counts(self, base, s3):
        spec = _spec(
            fmt="excel",
            location={"type": "s3", "uri": "s3://bkt/in/book.xlsx"},
            output={"location": {"type": "s3", "uri": "s3://bkt/xl/"}},
        )
        result = _run(spec, base, s3_client=s3[1])
        assert result.status == "succeeded", result.error
        assert [(a.path, a.rows) for a in result.artifacts] == [
            ("s3://bkt/xl/book_Sheet.parquet", None)
        ]

    @pytest.mark.parametrize("kind", ["preview", "validate_schema", "generate_schema"])
    def test_interactive_kinds_on_s3_csv(self, base, s3, kind):
        spec = _spec(kind=kind, location=S3_IN, schema=SCHEMA, output=None)
        result = _run(spec, base, s3_client=s3[1])
        assert result.status == "succeeded", result.error
        assert result.counts["total_rows"] == 3

    @pytest.mark.parametrize("kind", ["preview", "generate_schema"])
    def test_s3_workbook_is_copied_for_the_interactive_kinds(self, base, s3, kind):
        spec = _spec(
            kind=kind, fmt="excel", location={"type": "s3", "uri": "s3://bkt/in/book.xlsx"}
        )
        result = _run(spec, base, s3_client=s3[1])
        assert result.status == "succeeded", result.error
        assert sorted(p.name for p in base.iterdir()) == ["in", "out"]

    def test_missing_object_and_refused_access(self, base, s3):
        missing = _run(
            _spec(location={"type": "s3", "uri": "s3://bkt/in/none.csv"}), base, s3_client=s3[1]
        )
        assert missing.error.code == "INPUT_UNREADABLE"
        from botocore.exceptions import ClientError

        denied = ClientError({"Error": {"Code": "AccessDenied", "Message": "no"}}, "GetObject")
        with patch("forklift.engine.forklift_core.ForkliftCore.process_csv", side_effect=denied):
            result = _run(_spec(location=S3_IN), base, s3_client=s3[1])
        assert result.error.code == "PERMISSION_DENIED"


class TestErrorClassification:
    @pytest.mark.parametrize(
        "error, code, retryable",
        [
            (PermissionError("no"), "PERMISSION_DENIED", False),
            (UnicodeDecodeError("utf-8", b"\xe9", 0, 1, "bad"), "ENCODING_ERROR", False),
            (ConnectionResetError("reset"), "INPUT_UNREADABLE", True),
            (TimeoutError("slow"), "INPUT_UNREADABLE", True),
            (EOFError("short"), "INPUT_UNREADABLE", False),
            (
                pa.ArrowInvalid("CSV parse error: Expected 3 columns, got 4: 1,2,3,4"),
                "INPUT_UNREADABLE",
                False,
            ),
            (KeyError("name"), "INTERNAL", False),
            (ValueError("Sheet 'x' not found in workbook"), "INPUT_UNREADABLE", False),
            (ImportError("No module named 'openpyxl'"), "INTERNAL", False),
        ],
    )
    def test_codes_by_type(self, error, code, retryable):
        from forklift.jobs.errors import classify_error

        got, message, again = classify_error(error)
        assert (got, again) == (code, retryable)
        assert "1,2,3,4" not in message

    def test_s3_and_credential_errors(self):
        from botocore.exceptions import ClientError, NoCredentialsError

        from forklift.jobs.errors import classify_error

        def client_error(code):
            return ClientError({"Error": {"Code": code, "Message": "m"}}, "GetObject")

        assert classify_error(client_error("NoSuchKey"))[0] == "INPUT_UNREADABLE"
        assert classify_error(client_error("403"))[0] == "PERMISSION_DENIED"
        assert classify_error(client_error("SlowDown"))[2] is True
        assert classify_error(NoCredentialsError())[0] == "PERMISSION_DENIED"

    def test_messages_are_cut_short_and_lose_paths_and_secrets(self, tmp_path):
        from forklift.jobs.errors import MAX_MESSAGE_LENGTH, classify_error

        error = ValueError(f"{tmp_path}/in/x.csv and {tmp_path} with Pwd=hunter22;" + "y" * 5000)
        _, message, _ = classify_error(error, secrets=["Pwd=hunter22;", ""], base_dir=tmp_path)
        assert message.startswith("in/x.csv and . with ")
        assert "hunter22" not in message
        assert len(message) == MAX_MESSAGE_LENGTH and message.endswith("(cut short)")
        assert classify_error(KeyError("name"))[1] == "KeyError: 'name'"

    def test_database_error_without_a_sqlstate(self):
        from forklift.jobs.errors import classify_error

        class Error(Exception):
            pass

        Error.__module__ = "pyodbc.driver"
        assert classify_error(Error("x"))[1] == "Error"


def test_cell_text():
    import datetime
    import decimal

    from forklift.jobs.runner import _cell_text

    assert _cell_text(None) is None and _cell_text("a") == "a"
    assert _cell_text(1.5) == "1.5" and _cell_text(2.0) == "2" and _cell_text(1e20) == "1e+20"
    assert _cell_text(datetime.time(1, 2)) == "01:02:00"
    assert _cell_text(decimal.Decimal("1.10")) == "1.10"
