"""Tests for the WP2 review fixes: schema generation without pandas, PII opt-in, inference.

Covers metadata generation on ``pyarrow.compute``, ``include_value_statistics``, the CSV/Excel/
Parquet sampling rewrite, Arrow -> Forklift type strings, primary key inference rules and
special type detection.
"""

import datetime
import io
import json
import math
import os
import sys
from decimal import Decimal
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from forklift import api
from forklift.schema.csv_schema_importer import CsvSchemaImporter
from forklift.schema.excel_schema_importer import ExcelSchemaImporter
from forklift.schema.fwf.validation.parquet_types import ParquetTypeValidator
from forklift.schema.generator import inference as inference_module
from forklift.schema.generator.core import FileType, SchemaGenerationConfig, SchemaGenerator
from forklift.schema.generator.inference import DataTypeInferrer
from forklift.schema.processors import metadata as metadata_module
from forklift.schema.processors.config_parser import ConfigurationParser
from forklift.schema.processors.json_schema import JSONSchemaProcessor
from forklift.schema.processors.metadata import MetadataGenerator
from forklift.schema.types.special_types import SpecialTypeDetector
from forklift.schema.utils.formatters import SchemaFormatter
from forklift.schema.utils.helpers import (
    get_parquet_type_string,
    parquet_type_string_to_arrow,
    quantile_label,
    source_basename,
    split_name_tokens,
    to_json_safe,
    validate_quantiles,
)

VALUE_BEARING_KEYS = {
    "top_values",
    "bottom_values",
    "suggested_enum_values",
    "min_value",
    "max_value",
    "median",
    "quantiles",
    "range",
}


def strict_json_loads(text):
    """Parse JSON and fail on the non-standard NaN / Infinity constants."""

    def reject(constant):
        raise AssertionError(f"invalid JSON constant {constant}")

    return json.loads(text, parse_constant=reject)


def column_level_keys(metadata):
    """Keys of the per-column statistics and enum suggestions (not the config echo)."""
    return collect_keys(metadata["column_metadata"]) | collect_keys(metadata["enum_suggestions"])


def collect_keys(node, found=None):
    """All dict keys anywhere in a nested structure."""
    found = set() if found is None else found
    if isinstance(node, dict):
        for key, value in node.items():
            found.add(key)
            collect_keys(value, found)
    elif isinstance(node, list):
        for item in node:
            collect_keys(item, found)
    return found


class TrackingStream(io.BytesIO):
    """Binary stream that records how many bytes were read."""

    def __init__(self, data):
        super().__init__(data)
        self.bytes_read = 0

    def read(self, size=-1):
        chunk = super().read(size)
        self.bytes_read += len(chunk)
        return chunk

    def readinto(self, buffer):
        count = super().readinto(buffer)
        self.bytes_read += count
        return count


class ForwardOnlyStream:
    """A readable binary stream without random access (like an HTTP body)."""

    def __init__(self, data):
        self._buffer = io.BytesIO(data)
        self.closed = False

    def read(self, size=-1):
        return self._buffer.read(size)

    def readable(self):
        return True

    def seekable(self):
        return False

    def close(self):
        self.closed = True


class FakeS3:
    """Stand-in for ``UnifiedIOHandler.open_for_read`` returning a context manager."""

    def __init__(self, stream):
        self.stream = stream
        self.calls = []

    def open_for_read(self, path, encoding="utf-8", **kwargs):
        self.calls.append((path, encoding))
        return self

    def __enter__(self):
        return self.stream

    def __exit__(self, *exc):
        return False


@pytest.fixture
def pii_csv(tmp_path):
    path = tmp_path / "private_dir" / "people.csv"
    path.parent.mkdir()
    path.write_text(
        "id,name,ssn,email,status,score\n"
        "1,Ada Lovelace,123-45-6789,ada@example.com,active,10.5\n"
        "2,Grace Hopper,987-65-4321,grace@example.com,active,20.5\n"
        "3,Alan Turing,555-44-3333,alan@example.com,active,30.5\n"
        "4,Edsger Dijkstra,111-22-4444,edsger@example.com,active,40.5\n"
        "5,Barbara Liskov,222-33-5555,barbara@example.com,inactive,50.5\n"
        "6,Donald Knuth,333-44-6666,donald@example.com,active,60.5\n"
        "7,Margaret Hamilton,444-55-7777,margaret@example.com,active,70.5\n"
        "8,Ken Thompson,666-77-8888,ken@example.com,active,80.5\n"
        "9,Dennis Ritchie,777-88-9999,dennis@example.com,active,90.5\n"
        "10,Linus Torvalds,888-99-0000,linus@example.com,active,100.5\n"
    )
    return path


RAW_VALUES = [
    "123-45-6789",
    "ada@example.com",
    "Ada Lovelace",
    "Edsger Dijkstra",
    "888-99-0000",
]


# ---------------------------------------------------------------------------
# 1. No pandas in schema generation
# ---------------------------------------------------------------------------


def mixed_table():
    return pa.table(
        {
            "i": pa.array([1, 2, 3, None]),
            "f": pa.array([1.5, None, 3.5, 4.5]),
            "s": pa.array(["a", "b", "a", None]),
            "b": pa.array([True, False, True, None]),
            "lst": pa.array([[1, 2], None, [3], []]),
            "st": pa.array([{"x": 1}, None, {"x": 2}, {"x": 1}]),
            "d": pa.array(["a", "b", "a", "a"]).dictionary_encode(),
            "dt": pa.array([datetime.date(2020, 1, 1)] * 4),
            "dec": pa.array([Decimal("1.50")] * 4, pa.decimal128(10, 2)),
        }
    )


NO_PANDAS_SCRIPT = """
import sys

class BlockPandas:
    def find_spec(self, name, path=None, target=None):
        if name.split(".")[0] in ("pandas", "polars"):
            raise ImportError("blocked: " + name)

sys.meta_path.insert(0, BlockPandas())

import datetime
import json
from decimal import Decimal

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq

from forklift.schema.generator.inference import DataTypeInferrer
from forklift.schema.processors.config_parser import ConfigurationParser
from forklift.schema.processors.json_schema import JSONSchemaProcessor
from forklift.schema.processors.metadata import MetadataGenerator
from forklift.schema.utils.formatters import SchemaFormatter

table = pa.table({
    "i": pa.array([1, 2, 3, None]),
    "f": pa.array([1.5, None, 3.5, float("nan")]),
    "s": pa.array(["a", "b", "a", None]),
    "b": pa.array([True, False, True, None]),
    "lst": pa.array([[1, 2], None, [3], []]),
    "st": pa.array([{"x": 1}, None, {"x": 2}, {"x": 1}]),
    "d": pa.array(["a", "b", "a", "a"]).dictionary_encode(),
    "dt": pa.array([datetime.date(2020, 1, 1)] * 4),
    "dec": pa.array([Decimal("1.50")] * 4, pa.decimal128(10, 2)),
})
for flag in (False, True):
    MetadataGenerator().generate_metadata(table, {"include_value_statistics": flag})
JSONSchemaProcessor().generate_sample_data(table)
ConfigurationParser()._infer_primary_key_from_metadata(table)
SchemaFormatter.format_schema_json({"x": float("nan")})

workdir = sys.argv[1]
csv_path = workdir + "/a.csv"
open(csv_path, "w").write("a,b\\n1,x\\n2,y\\n")
inferrer = DataTypeInferrer()
assert inferrer.read_csv_sample(csv_path, 1000).num_rows == 2
assert inferrer.read_csv_sample(csv_path, None).num_rows == 2

xlsx_path = workdir + "/a.xlsx"
book = openpyxl.Workbook()
book.active.append(["a", "b"])
book.active.append([1, "x"])
book.save(xlsx_path)
assert inferrer.read_excel_sample(xlsx_path, 10).num_rows == 1

parquet_path = workdir + "/a.parquet"
pq.write_table(pa.table({"a": [1, 2, 3]}), parquet_path)
assert inferrer.read_parquet_sample(parquet_path, 2).num_rows == 2
assert inferrer.read_parquet_sample(parquet_path, None).num_rows == 3

assert "pandas" not in sys.modules and "polars" not in sys.modules
print("OK")
"""


class TestPandasFree:
    def test_schema_components_run_without_pandas_installed(self, tmp_path):
        import subprocess

        import forklift

        script = tmp_path / "no_pandas_check.py"
        script.write_text(NO_PANDAS_SCRIPT)
        src_dir = str(Path(forklift.__file__).resolve().parent.parent)
        env = dict(os.environ, PYTHONPATH=src_dir)
        result = subprocess.run(
            [sys.executable, str(script), str(tmp_path)],
            capture_output=True,
            text=True,
            env=env,
            cwd=str(tmp_path),
        )
        assert result.returncode == 0, result.stderr[-2000:]
        assert result.stdout.strip().endswith("OK")

    def test_no_pandas_references_in_owned_modules(self):
        import forklift.schema.generator.inference as inference
        import forklift.schema.processors.json_schema as json_schema
        import forklift.schema.processors.metadata as metadata
        import forklift.schema.types.special_types as special_types
        import forklift.schema.utils.helpers as helpers

        for module in (inference, json_schema, metadata, special_types, helpers):
            source = Path(module.__file__).read_text()
            assert "import pandas" not in source
            assert "from pandas" not in source
            assert "to_pandas" not in source
            assert "from_pandas" not in source


# ---------------------------------------------------------------------------
# 1/3. Metadata on pyarrow.compute: correctness and robustness
# ---------------------------------------------------------------------------


def stats_for(values, **config):
    table = pa.table({"v": values})
    config.setdefault("include_value_statistics", True)
    return MetadataGenerator().generate_metadata(table, config)["column_metadata"]["v"]


class TestMetadataStatistics:
    def test_numeric_statistics_match_reference(self):
        import statistics

        data = [1, 2, 3, 4, 100, 7, 9]
        meta = stats_for(pa.array(data))
        assert meta["min_value"] == 1.0
        assert meta["max_value"] == 100.0
        assert meta["mean"] == pytest.approx(statistics.mean(data))
        assert meta["median"] == pytest.approx(statistics.median(data))
        assert meta["std_dev"] == pytest.approx(statistics.stdev(data))
        assert meta["variance"] == pytest.approx(statistics.variance(data))
        assert meta["range"] == 99.0
        assert meta["quantiles"]["quantile_50"] == pytest.approx(statistics.median(data))
        # IQR rule: 100 is an outlier, nothing else is
        assert meta["outlier_count"] == 1
        assert meta["outlier_percentage"] == pytest.approx(100 / 7)
        assert meta["distinct_count"] == 7
        assert meta["uniqueness_ratio"] == 1.0

    def test_single_row_has_no_nan_deviation(self):
        meta = stats_for(pa.array([5.0]))
        assert meta["std_dev"] is None
        assert meta["variance"] is None
        assert meta["coefficient_of_variation"] is None
        assert meta["mean"] == 5.0
        json.dumps(meta, allow_nan=False)

    def test_nan_and_infinity_never_reach_the_output(self):
        table = pa.table(
            {
                "f": pa.array([1.0, float("nan"), float("inf"), None, -float("inf")]),
                "n": pa.array([float("nan")] * 5),
            }
        )
        metadata = MetadataGenerator().generate_metadata(table, {"include_value_statistics": True})
        text = json.dumps(metadata, allow_nan=False)
        strict_json_loads(text)
        assert metadata["column_metadata"]["f"]["nan_count"] == 1
        assert metadata["column_metadata"]["f"]["null_count"] == 1
        # a column of only NaN has no analysable values
        assert "mean" not in metadata["column_metadata"]["n"]

    def test_schema_output_is_strict_json_with_nan_data(self, tmp_path):
        path = tmp_path / "nan.parquet"
        pq.write_table(
            pa.table({"x": pa.array([1.0, float("nan"), float("inf")]), "y": [1, 2, 3]}), path
        )
        config = SchemaGenerationConfig(
            input_path=str(path),
            file_type=FileType.PARQUET,
            include_sample_data=True,
            include_value_statistics=True,
        )
        generator = SchemaGenerator(config)
        schema = generator.generate_schema()
        strict_json_loads(generator.formatter.format_schema_json(schema))
        assert schema["x-sample"]["rows"][1]["x"] is None

    def test_quantile_labels_use_proper_rounding(self):
        meta = stats_for(pa.array(range(1, 101)), quantiles=[0.29, 0.07, 0.995, 0.5, 1, 0])
        assert list(meta["quantiles"]) == [
            "quantile_29",
            "quantile_7",
            "quantile_99_5",
            "quantile_50",
            "quantile_100",
            "quantile_0",
        ]
        assert meta["quantiles"]["quantile_29"] == pytest.approx(29.71)
        assert meta["quantiles"]["quantile_100"] == 100.0

    @pytest.mark.parametrize(
        "q,label", [(0.29, "29"), (0.57, "57"), (0.58, "58"), (0.1, "10"), (0.025, "2_5")]
    )
    def test_quantile_label_is_exact(self, q, label):
        assert quantile_label(q) == label
        # the old int(q * 100) was off by one for some of these
        assert int(0.29 * 100) == 28

    @pytest.mark.parametrize("bad", [-0.1, 1.5, float("nan"), float("inf"), "0.5", True])
    def test_invalid_quantiles_raise_clear_error(self, bad):
        with pytest.raises(ValueError, match="quantile"):
            validate_quantiles([0.25, bad])
        with pytest.raises(ValueError, match="quantile"):
            SchemaGenerationConfig(input_path="x.csv", file_type=FileType.CSV, quantiles=[bad])
        with pytest.raises(ValueError, match="quantile"):
            MetadataGenerator().generate_metadata(pa.table({"v": [1, 2]}), {"quantiles": [bad]})

    def test_missing_quantiles_select_the_defaults(self):
        assert validate_quantiles(None) == [0.25, 0.5, 0.75, 0.9, 0.95, 0.99]
        config = SchemaGenerationConfig(input_path="x.csv", file_type=FileType.CSV, quantiles=None)
        assert config.quantiles == [0.25, 0.5, 0.75, 0.9, 0.95, 0.99]

    def test_quantile_boundaries_are_valid(self):
        assert validate_quantiles([0, 1, 0.5]) == [0.0, 1.0, 0.5]
        assert SchemaGenerationConfig(
            input_path="x.csv", file_type=FileType.CSV, quantiles=[0, 1]
        ).quantiles == [0.0, 1.0]

    def test_top_and_bottom_values_ordering(self):
        values = ["a"] * 5 + ["b"] * 3 + ["c"] * 3 + ["d", "e", "f"]
        meta = stats_for(pa.array(values), top_n_values=2)
        assert [(v["value"], v["count"]) for v in meta["top_values"]] == [("a", 5), ("b", 3)]
        # ties keep first-seen order; bottom values are the least frequent, least last
        assert [(v["value"], v["count"]) for v in meta["bottom_values"]] == [("e", 1), ("f", 1)]
        assert meta["top_values"][0]["percentage"] == pytest.approx(5 / 14 * 100)

    def test_unhashable_columns_are_skipped_not_crashed(self):
        table = mixed_table()
        metadata = MetadataGenerator().generate_metadata(table, {"include_value_statistics": True})
        for name in ("lst", "st"):
            column = metadata["column_metadata"][name]
            assert "distinct_count" not in column
            assert "top_values" not in column
            assert column["null_count"] == 1
        map_table = pa.table(
            {"m": pa.array([[("k", 1)], None], type=pa.map_(pa.string(), pa.int64()))}
        )
        meta = MetadataGenerator().generate_metadata(map_table, {})
        assert meta["column_metadata"]["m"]["null_count"] == 1

    def test_dictionary_columns_use_decoded_values(self):
        meta = stats_for(pa.array(["a", "b", "a", "a"]).dictionary_encode())
        assert meta["distinct_count"] == 2
        assert meta["top_values"][0] == {"value": "a", "count": 3, "percentage": 75.0}
        assert meta["max_length"] == 1  # string statistics work on decoded values

    def test_string_statistics(self):
        meta = stats_for(pa.array(["Hello World", "UPPER", "lower", None, "", "caf\u00e9 1!"]))
        assert meta["min_length"] == 0
        assert meta["max_length"] == 11
        assert meta["empty_strings"] == 2  # one "" plus one null
        assert meta["contains_whitespace"] == 2
        assert meta["contains_numbers"] == 1
        assert meta["contains_special_chars"] == 1  # "caf\u00e9 1!" ('!' and '\u00e9')
        assert meta["all_uppercase"] == 1
        assert meta["all_lowercase"] == 2
        assert meta["ascii_only"] == 4
        assert meta["non_ascii_count"] == 1

    def test_boolean_statistics(self):
        meta = stats_for(pa.array([True, True, False, None]))
        assert meta["true_count"] == 2
        assert meta["false_count"] == 1
        assert meta["true_percentage"] == pytest.approx(200 / 3)

    def test_statistics_errors_do_not_leak_values(self, monkeypatch):
        def boom(*args, **kwargs):
            raise RuntimeError("secret-cell-value")

        monkeypatch.setattr(metadata_module.pc, "mean", boom)
        meta = stats_for(pa.array([1.0, 2.0]))
        assert "error" in meta
        assert "secret-cell-value" not in json.dumps(meta)

    def test_uses_long_integers_without_overflow(self):
        meta = stats_for(pa.array([2**62, 2**62, 2**62], pa.int64()))
        assert meta["mean"] == pytest.approx(2.0**62)
        assert math.isfinite(meta["variance"])


# ---------------------------------------------------------------------------
# 2. PII: value statistics are opt-in
# ---------------------------------------------------------------------------


def generate(path, **overrides):
    config = SchemaGenerationConfig(input_path=str(path), file_type=FileType.CSV, **overrides)
    generator = SchemaGenerator(config)
    return generator, generator.generate_schema()


class TestValueStatisticsOptIn:
    def test_default_schema_contains_no_cell_values(self, pii_csv):
        generator, schema = generate(pii_csv, enum_threshold=0.9)
        text = generator.formatter.format_schema_json(schema)
        for raw in RAW_VALUES:
            assert raw not in text
        assert str(pii_csv.parent) not in text
        assert not VALUE_BEARING_KEYS & column_level_keys(schema["x-metadata"])
        assert "x-sample" not in schema

    def test_default_keeps_counts_nulls_types_and_lengths(self, pii_csv):
        _, schema = generate(pii_csv, enum_threshold=0.9)
        columns = schema["x-metadata"]["column_metadata"]
        assert columns["ssn"]["distinct_count"] == 10
        assert columns["ssn"]["null_count"] == 0
        assert columns["ssn"]["parquet_type"] == "string"
        assert columns["ssn"]["min_length"] == 11
        assert columns["ssn"]["max_length"] == 11
        assert columns["score"]["mean"] == pytest.approx(55.5)
        assert columns["score"]["std_dev"] is not None
        assert "min_value" not in columns["score"]
        assert "quantiles" not in columns["score"]

    def test_enum_suggestion_has_no_value_lists(self, pii_csv):
        _, schema = generate(pii_csv, enum_threshold=0.5)
        suggestion = schema["x-metadata"]["enum_suggestions"]["status"]
        assert suggestion["is_enum_candidate"] is True
        assert suggestion["distinct_count"] == 2
        assert "suggested_enum_values" not in suggestion
        assert "active" not in suggestion["recommendation"]
        assert "inactive" not in suggestion["recommendation"]

    def test_opt_in_adds_value_statistics(self, pii_csv):
        _, schema = generate(pii_csv, enum_threshold=0.5, include_value_statistics=True)
        columns = schema["x-metadata"]["column_metadata"]
        assert columns["status"]["top_values"][0]["value"] == "active"
        assert columns["score"]["min_value"] == 10.5
        assert columns["score"]["max_value"] == 100.5
        assert "quantile_50" in columns["score"]["quantiles"]
        suggestion = schema["x-metadata"]["enum_suggestions"]["status"]
        assert suggestion["suggested_enum_values"] == ["active", "inactive"]
        assert "active" in suggestion["recommendation"]
        assert schema["x-metadata"]["analysis_config"]["include_value_statistics"] is True

    def test_source_file_is_a_basename(self, pii_csv):
        _, schema = generate(pii_csv)
        assert schema["x-generation"]["source_file"] == "people.csv"
        assert schema["x-metadata"]["table_metadata"]["source_file"] == "people.csv"

    @pytest.mark.parametrize(
        "path,expected",
        [
            ("/home/alice/data/a.csv", "a.csv"),
            ("s3://bucket/some/prefix/b.csv", "b.csv"),
            ("C:\\Users\\alice\\c.csv", "c.csv"),
            ("rel/d.csv", "d.csv"),
            ("", "unknown"),
        ],
    )
    def test_source_basename(self, path, expected):
        assert source_basename(path) == expected

    def test_metadata_generator_trims_source_path_itself(self):
        metadata = MetadataGenerator().generate_metadata(
            pa.table({"a": [1]}), {"source_file": "/secret/dir/file.csv"}
        )
        assert metadata["table_metadata"]["source_file"] == "file.csv"

    def test_separate_metadata_file_has_no_values_by_default(self, pii_csv, tmp_path):
        out = tmp_path / "meta.json"
        config = SchemaGenerationConfig(
            input_path=str(pii_csv), file_type=FileType.CSV, metadata_output_path=str(out)
        )
        generator = SchemaGenerator(config)
        generator.generate_and_save_metadata(generator._read_sample_data())
        text = out.read_text()
        for raw in RAW_VALUES:
            assert raw not in text
        assert str(pii_csv.parent) not in text
        assert not VALUE_BEARING_KEYS & column_level_keys(json.loads(text))

    def test_sample_data_keeps_its_own_opt_in(self, pii_csv):
        _, schema = generate(pii_csv, include_sample_data=True)
        assert schema["x-sample"]["rows"][0]["name"] == "Ada Lovelace"
        # value statistics stay off
        assert not VALUE_BEARING_KEYS & column_level_keys(schema["x-metadata"])

    def test_config_default_is_off(self):
        config = SchemaGenerationConfig(input_path="x.csv", file_type=FileType.CSV)
        assert config.include_value_statistics is False

    def test_primary_key_inference_does_not_embed_values(self, pii_csv):
        _, schema = generate(pii_csv, infer_primary_key_from_metadata=True)
        assert schema["x-primaryKey"]["columns"] == ["id"]
        assert "1" not in json.dumps(
            schema["x-primaryKey"]["inference_metadata"]["alternative_candidates"]
        )


class TestApiPlumbing:
    def test_csv_api_defaults_to_no_value_statistics(self, pii_csv):
        schema = api.generate_schema_from_csv(pii_csv)
        for raw in RAW_VALUES:
            assert raw not in json.dumps(schema)
        assert not VALUE_BEARING_KEYS & column_level_keys(schema["x-metadata"])

    def test_csv_api_opt_in(self, pii_csv):
        schema = api.generate_schema_from_csv(pii_csv, include_value_statistics=True)
        assert schema["x-metadata"]["column_metadata"]["score"]["max_value"] == 100.5

    def test_excel_api_option(self, tmp_path):
        path = tmp_path / "e.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(["name", "score"])
        workbook.active.append(["Ada", 1.5])
        workbook.active.append(["Grace", 2.5])
        workbook.save(path)
        default = api.generate_schema_from_excel(path)
        assert "Ada" not in json.dumps(default)
        opted_in = api.generate_schema_from_excel(path, include_value_statistics=True)
        assert opted_in["x-metadata"]["column_metadata"]["score"]["max_value"] == 2.5

    def test_parquet_api_option(self, tmp_path):
        path = tmp_path / "p.parquet"
        pq.write_table(pa.table({"name": ["Ada", "Grace"], "score": [1.5, 2.5]}), path)
        default = api.generate_schema_from_parquet(path)
        assert "Ada" not in json.dumps(default)
        opted_in = api.generate_schema_from_parquet(path, include_value_statistics=True)
        assert opted_in["x-metadata"]["column_metadata"]["name"]["top_values"]

    def test_generate_and_save_schema_accepts_option(self, pii_csv, tmp_path):
        out = tmp_path / "schema.json"
        api.generate_and_save_schema(pii_csv, out, "csv", include_value_statistics=True)
        saved = json.loads(out.read_text())
        assert "top_values" in column_level_keys(saved["x-metadata"])
        out2 = tmp_path / "schema2.json"
        api.generate_and_save_schema(pii_csv, out2, "csv")
        assert "top_values" not in column_level_keys(json.loads(out2.read_text())["x-metadata"])


class TestJsonSafeSamples:
    def test_sample_rows_are_json_safe(self):
        table = pa.table(
            {
                "day": pa.array([datetime.date(2020, 1, 2)]),
                "ts": pa.array([datetime.datetime(2020, 1, 2, 3, 4, 5)]),
                "zts": pa.array(
                    [datetime.datetime(2020, 1, 2, 3, 4, 5, tzinfo=datetime.timezone.utc)],
                    pa.timestamp("us", tz="UTC"),
                ),
                "dec": pa.array([Decimal("12.3400")], pa.decimal128(10, 4)),
                "f": pa.array([float("nan")]),
                "inf": pa.array([float("inf")]),
                "bin": pa.array([b"\x00\x01"]),
                "lst": pa.array([[1, 2]]),
                "t": pa.array([datetime.time(1, 2, 3)]),
            }
        )
        rows = JSONSchemaProcessor().generate_sample_data(table)["rows"]
        assert rows == [
            {
                "day": "2020-01-02",
                "ts": "2020-01-02T03:04:05",
                "zts": "2020-01-02T03:04:05+00:00",
                "dec": "12.3400",
                "f": None,
                "inf": None,
                "bin": "AAE=",
                "lst": [1, 2],
                "t": "01:02:03",
            }
        ]
        strict_json_loads(json.dumps(rows, allow_nan=False))

    def test_to_json_safe_conversions(self):
        assert to_json_safe(float("-inf")) is None
        assert to_json_safe(Decimal("NaN")) is None
        assert to_json_safe({1: (1, 2)}) == {"1": [1, 2]}
        assert to_json_safe(datetime.timedelta(seconds=90)) == "0:01:30"

    def test_formatter_never_writes_bare_nan(self):
        text = SchemaFormatter.format_schema_json(
            {"a": float("nan"), "b": [float("inf")], "c": Decimal("1.5")}
        )
        assert strict_json_loads(text) == {"a": None, "b": [None], "c": "1.5"}


# ---------------------------------------------------------------------------
# 4. Inference sampling
# ---------------------------------------------------------------------------


def types_of(table):
    return {field.name: field.type for field in table.schema}


class TestCsvSampling:
    def test_leading_zeros_and_na_are_preserved(self, tmp_path):
        path = tmp_path / "z.csv"
        path.write_text("zip,code,country\n02134,00123,NA\n90210,00456,US\n")
        for nrows in (1000, None):
            table = DataTypeInferrer().read_csv_sample(path, nrows)
            assert table.column("zip").to_pylist() == ["02134", "90210"]
            assert table.column("code").to_pylist() == ["00123", "00456"]
            assert table.column("country").to_pylist() == ["NA", "US"]
            assert set(types_of(table).values()) == {pa.string()}

    def test_nrows_and_unlimited_agree_for_the_same_data(self, tmp_path):
        path = tmp_path / "t.csv"
        path.write_text(
            "i,f,b,s,d,ts,tz,nul\n"
            "1,1.5,true,x,2020-01-01,2020-01-01 10:00:00,2020-01-01T10:00:00Z,\n"
            "2,2,FALSE,y,2020-01-02,2020-01-02 11:00:00,2020-01-02T10:00:00+01:00,NULL\n"
        )
        limited = DataTypeInferrer().read_csv_sample(path, 1000)
        unlimited = DataTypeInferrer().read_csv_sample(path, None)
        assert limited.equals(unlimited)
        assert types_of(limited) == {
            "i": pa.int64(),
            "f": pa.float64(),
            "b": pa.bool_(),
            "s": pa.string(),
            "d": pa.date32(),
            "ts": pa.timestamp("s"),
            "tz": pa.timestamp("s", tz="UTC"),
            "nul": pa.string(),
        }

    def test_strings_are_plain_not_dictionary_encoded(self, tmp_path):
        path = tmp_path / "s.csv"
        path.write_text("c\n" + "\n".join(["a", "b"] * 50) + "\n")
        table = DataTypeInferrer().read_csv_sample(path, None)
        assert table.schema.field("c").type == pa.string()

    @pytest.mark.parametrize(
        "values,expected",
        [
            (["1", "-2", "0"], pa.int64()),
            (["007", "1"], pa.string()),
            (["00", "1"], pa.string()),
            (["+1", "2"], pa.string()),
            (["1.5", "2"], pa.float64()),
            (["1e3", "-2.5E-2", ".5", "1."], pa.float64()),
            (["01.5", "2.5"], pa.string()),
            (["1,234", "2"], pa.string()),
            (["true", "False", "TRUE"], pa.bool_()),
            (["yes", "no"], pa.string()),
            (["1", "0"], pa.int64()),
            (["2020-01-01", "1999-12-31"], pa.date32()),
            (["2020-02-30", "2020-01-01"], pa.string()),
            (["2020-1-1"], pa.string()),
            (["2020-01-01", "2020-01-01 10:00:00"], pa.string()),
            (["2020-01-01 10:00:00.123", "2020-01-01 10:00:00"], pa.timestamp("ms")),
            (["2020-01-01T10:00:00.123456"], pa.timestamp("us")),
            (["2020-01-01T10:00:00.123456789"], pa.timestamp("ns")),
            (["2020-01-01 10:00"], pa.timestamp("s")),
            (["2020-01-01T10:00:00Z", "2020-01-01T10:00:00+05:30"], pa.timestamp("s", tz="UTC")),
            (["2020-01-01T10:00:00Z", "2020-01-01T10:00:00"], pa.string()),
            (["9223372036854775807"], pa.int64()),
            (["18446744073709551615"], pa.uint64()),
            (["99999999999999999999999"], pa.string()),
            (["-9223372036854775809"], pa.string()),
            ([" 12"], pa.string()),
        ],
    )
    def test_type_rules(self, values, expected):
        table = pa.table({"c": pa.array(values, pa.string())})
        inferred = DataTypeInferrer().infer_types_from_strings(table)
        assert inferred.schema.field("c").type == expected

    def test_null_tokens(self):
        table = pa.table(
            {
                "c": pa.array(["1", "", "NULL", "null", "N/A", "n/a", "#N/A", "NaN", "nan", "2"]),
                "na": pa.array(["NA"] * 10),
                "all_null": pa.array(["", "NULL"] * 5),
            }
        )
        inferred = DataTypeInferrer().infer_types_from_strings(table)
        assert inferred.column("c").to_pylist() == [
            1,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            2,
        ]
        assert inferred.column("na").to_pylist() == ["NA"] * 10  # "NA" is a value, not a null
        assert inferred.column("all_null").to_pylist() == [None] * 10
        assert inferred.schema.field("all_null").type == pa.string()
        custom = DataTypeInferrer().infer_types_from_strings(
            pa.table({"c": pa.array(["1", "NA"])}), null_tokens=["NA"]
        )
        assert custom.column("c").to_pylist() == [1, None]

    def test_nrows_limits_rows_and_headers_become_columns(self, tmp_path):
        path = tmp_path / "n.csv"
        path.write_text("a,b\n" + "".join(f"{i},x{i}\n" for i in range(100)))
        table = DataTypeInferrer().read_csv_sample(path, 7)
        assert table.num_rows == 7
        assert table.column("a").to_pylist() == list(range(7))

    def test_header_only_files(self, tmp_path):
        for text in ("id,name,value", "id,name,value\n", "id,name,value\r\n"):
            path = tmp_path / "h.csv"
            path.write_text(text)
            table = DataTypeInferrer().read_csv_sample(path, 10)
            assert table.num_rows == 0
            assert table.column_names == ["id", "name", "value"]

    def test_empty_file_raises(self, tmp_path):
        path = tmp_path / "empty.csv"
        path.write_text("")
        with pytest.raises(ValueError, match="empty"):
            DataTypeInferrer().read_csv_sample(path, 10)

    def test_bad_rows_raise_without_echoing_data(self, tmp_path):
        path = tmp_path / "bad.csv"
        path.write_text("a,b\n1,2\n3,4,top-secret-value\n")
        with pytest.raises(ValueError) as excinfo:
            DataTypeInferrer().read_csv_sample(path, None)
        assert "top-secret-value" not in str(excinfo.value)
        assert "row 3" in str(excinfo.value)
        assert excinfo.value.__cause__ is None
        assert excinfo.value.__suppress_context__ is True

    def test_invalid_encoding_data_raises_without_echoing_data(self, tmp_path):
        path = tmp_path / "latin.csv"
        path.write_bytes("a\nsecr\u00e9t\n".encode("latin-1"))
        with pytest.raises(ValueError, match="encoding") as excinfo:
            DataTypeInferrer().read_csv_sample(path, None)
        assert "secr" not in str(excinfo.value)

    def test_configured_encoding_is_used(self, tmp_path):
        path = tmp_path / "latin.csv"
        path.write_bytes("name;n\ncaf\u00e9;1\n".encode("latin-1"))
        table = DataTypeInferrer().read_csv_sample(path, 10, delimiter=";", encoding="latin-1")
        assert table.column("name").to_pylist() == ["caf\u00e9"]
        assert table.column("n").to_pylist() == [1]

    def test_unknown_encoding(self, tmp_path):
        path = tmp_path / "a.csv"
        path.write_text("a\n1\n")
        with pytest.raises(ValueError, match="encoding"):
            DataTypeInferrer().read_csv_sample(path, 10, encoding="no-such-encoding")

    def test_utf8_bom_is_not_part_of_the_first_column_name(self, tmp_path):
        path = tmp_path / "bom.csv"
        path.write_bytes(b"\xef\xbb\xbfid,name\n1,a\n")
        assert DataTypeInferrer().read_csv_sample(path, 10).column_names == ["id", "name"]
        assert DataTypeInferrer().read_csv_sample(path, 10, encoding="utf-8-sig").column_names == [
            "id",
            "name",
        ]

    def test_blank_and_duplicate_header_names(self, tmp_path):
        path = tmp_path / "dup.csv"
        path.write_text(",a,a,a.1\n1,2,3,4\n")
        table = DataTypeInferrer().read_csv_sample(path, 10)
        assert table.column_names == ["Unnamed: 0", "a", "a.1", "a.1.1"]

    def test_quoted_newlines_are_one_row(self, tmp_path):
        path = tmp_path / "q.csv"
        path.write_text('id,text\n1,"line one\nline two"\n2,"x"\n3,y\n')
        table = DataTypeInferrer().read_csv_sample(path, 2)
        assert table.num_rows == 2
        assert table.column("text").to_pylist() == ["line one\nline two", "x"]

    @pytest.mark.parametrize(
        "nrows", [0, -1, True, 1.5, "10"], ids=["zero", "negative", "bool", "float", "str"]
    )
    def test_invalid_nrows(self, tmp_path, nrows):
        path = tmp_path / "a.csv"
        path.write_text("a\n1\n")
        with pytest.raises(ValueError, match="nrows"):
            DataTypeInferrer().read_csv_sample(path, nrows)

    def test_sampling_works_without_default_column_type(self, tmp_path, monkeypatch):
        monkeypatch.setattr(inference_module, "_SUPPORTS_DEFAULT_COLUMN_TYPE", False)
        path = tmp_path / "old.csv"
        path.write_text("zip,n\n02134,1\n00123,2\n")
        table = DataTypeInferrer().read_csv_sample(path, 5)
        assert table.column("zip").to_pylist() == ["02134", "00123"]
        assert table.column("n").to_pylist() == [1, 2]


class TestS3CsvSampling:
    def inferrer_with(self, stream):
        inferrer = DataTypeInferrer()
        inferrer.io_handler = FakeS3(stream)
        return inferrer

    def test_s3_reads_binary_and_stops_early(self):
        rows = "".join(f"{i},name{i},{i * 1.5}\n" for i in range(300_000))
        data = ("id,name,value\n" + rows).encode()
        assert len(data) > 7_000_000
        stream = TrackingStream(data)
        inferrer = self.inferrer_with(stream)
        table = inferrer.read_csv_sample("s3://bucket/big.csv", 10)
        assert inferrer.io_handler.calls == [("s3://bucket/big.csv", "binary")]
        assert table.num_rows == 10
        assert table.schema.field("value").type == pa.float64()
        assert stream.bytes_read < 4_000_000  # Arrow's read-ahead only, not the whole 7+ MB

    def test_s3_honours_encoding(self):
        stream = TrackingStream("a,b\ncaf\u00e9,1\n".encode("latin-1"))
        table = self.inferrer_with(stream).read_csv_sample(
            "s3://bucket/k.csv", 10, encoding="latin-1"
        )
        assert table.column("a").to_pylist() == ["caf\u00e9"]

    def test_s3_counts_records_not_physical_lines(self):
        stream = TrackingStream(b'id,text\n1,"a\nb\nc"\n2,"d\ne"\n3,f\n4,g\n')
        table = self.inferrer_with(stream).read_csv_sample("s3://bucket/k.csv", 3)
        assert table.num_rows == 3
        assert table.column("text").to_pylist() == ["a\nb\nc", "d\ne", "f"]

    def test_s3_forward_only_stream(self):
        stream = ForwardOnlyStream(b"id,name\n1,a\n2,b\n")
        table = self.inferrer_with(stream).read_csv_sample("s3://bucket/k.csv", None)
        assert table.num_rows == 2

    def test_s3_header_only_without_newline(self):
        class Reopening(FakeS3):
            def __enter__(self):
                return TrackingStream(b"id,name")

        inferrer = DataTypeInferrer()
        inferrer.io_handler = Reopening(None)
        table = inferrer.read_csv_sample("s3://bucket/k.csv", 5)
        assert table.column_names == ["id", "name"]
        assert table.num_rows == 0


class TestExcelSampling:
    @pytest.fixture
    def workbook_path(self, tmp_path):
        path = tmp_path / "book.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "First"
        sheet.append(["id", "code", "price", "when", "mixed", "empty"])
        for i in range(1, 51):
            sheet.append(
                [
                    i,
                    "0%04d" % i,
                    i * 1.5,
                    datetime.datetime(2020, 1, 1) + datetime.timedelta(days=i),
                    i if i % 2 else "text",
                    None,
                ]
            )
        sheet.append([None] * 6)  # trailing blank row
        second = workbook.create_sheet("Second")
        second.append(["only"])
        second.append([True])
        workbook.save(path)
        return path

    def test_types_and_nrows(self, workbook_path):
        table = DataTypeInferrer().read_excel_sample(workbook_path, 5)
        assert table.num_rows == 5
        assert types_of(table) == {
            "id": pa.int64(),
            "code": pa.string(),
            "price": pa.float64(),
            "when": pa.timestamp("us"),
            "mixed": pa.string(),
            "empty": pa.string(),
        }
        assert table.column("code").to_pylist()[0] == "00001"  # text cell keeps leading zeros
        assert table.column("mixed").to_pylist()[:2] == ["1", "text"]

    def test_all_rows_and_trailing_blank_rows(self, workbook_path):
        assert DataTypeInferrer().read_excel_sample(workbook_path, None).num_rows == 50

    def test_uses_read_only_openpyxl_and_closes_workbook(self, workbook_path, monkeypatch):
        seen = {}
        real_load = openpyxl.load_workbook

        def spy(*args, **kwargs):
            seen.update(kwargs)
            workbook = real_load(*args, **kwargs)
            real_close = workbook.close

            def close():
                seen["closed"] = True
                real_close()

            workbook.close = close
            return workbook

        monkeypatch.setattr(openpyxl, "load_workbook", spy)
        DataTypeInferrer().read_excel_sample(workbook_path, 3)
        assert seen["read_only"] is True
        assert seen["data_only"] is True
        assert seen["closed"] is True

    def test_only_needed_rows_are_requested(self, workbook_path, monkeypatch):
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet

        requested = {}
        real_iter_rows = ReadOnlyWorksheet.iter_rows

        def spy(self, *args, **kwargs):
            requested.update(kwargs)
            return real_iter_rows(self, *args, **kwargs)

        monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", spy)
        DataTypeInferrer().read_excel_sample(workbook_path, 4)
        assert requested["max_row"] == 5  # header + 4 data rows
        assert requested["values_only"] is True

    def test_sheet_selection(self, workbook_path):
        inferrer = DataTypeInferrer()
        assert inferrer.read_excel_sample(workbook_path, 10, "Second").column_names == ["only"]
        assert inferrer.read_excel_sample(workbook_path, 10, 1).column_names == ["only"]
        assert inferrer.read_excel_sample(workbook_path, 10, 0).column_names[0] == "id"
        with pytest.raises(ValueError, match="not found"):
            inferrer.read_excel_sample(workbook_path, 10, "Missing")
        with pytest.raises(ValueError, match="out of range"):
            inferrer.read_excel_sample(workbook_path, 10, 7)

    def test_wrong_dimension_metadata_does_not_truncate_the_sample(self, tmp_path):
        import re
        import zipfile

        good = tmp_path / "good.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(["a", "b", "c"])
        for i in range(5):
            workbook.active.append([i, i * 2, i * 3])
        workbook.save(good)

        bad = tmp_path / "bad.xlsx"
        with zipfile.ZipFile(good) as source, zipfile.ZipFile(bad, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    data = re.sub(rb"<dimension[^>]*/>", b'<dimension ref="A1"/>', data)
                target.writestr(item, data)

        for nrows in (None, 3):
            table = DataTypeInferrer().read_excel_sample(bad, nrows)
            assert table.column_names == ["a", "b", "c"]
            assert table.num_rows == (5 if nrows is None else 3)

    def test_legacy_xls_is_rejected_clearly(self, tmp_path):
        path = tmp_path / "old.xls"
        path.write_bytes(b"\xd0\xcf\x11\xe0")
        with pytest.raises(ValueError, match="xls"):
            DataTypeInferrer().read_excel_sample(path, 10)

    def test_s3_excel_is_read_from_a_binary_stream(self, workbook_path):
        inferrer = DataTypeInferrer()
        inferrer.io_handler = FakeS3(ForwardOnlyStream(workbook_path.read_bytes()))
        table = inferrer.read_excel_sample("s3://bucket/book.xlsx", 3)
        assert inferrer.io_handler.calls == [("s3://bucket/book.xlsx", "binary")]
        assert table.num_rows == 3

    def test_empty_workbook(self, tmp_path):
        path = tmp_path / "blank.xlsx"
        openpyxl.Workbook().save(path)
        assert DataTypeInferrer().read_excel_sample(path, 10).num_columns == 0


class TestParquetSampling:
    @pytest.fixture
    def parquet_path(self, tmp_path):
        path = tmp_path / "groups.parquet"
        pq.write_table(
            pa.table({"a": list(range(100)), "b": [f"x{i}" for i in range(100)]}),
            path,
            row_group_size=10,
        )
        return path

    def test_reads_only_the_needed_batches(self, parquet_path, monkeypatch):
        yielded = []
        real_file = pq.ParquetFile

        class Spy(real_file):
            def iter_batches(self, *args, **kwargs):
                for batch in super().iter_batches(*args, **kwargs):
                    yielded.append(batch.num_rows)
                    yield batch

        monkeypatch.setattr(pq, "ParquetFile", Spy)
        table = DataTypeInferrer().read_parquet_sample(parquet_path, 15)
        assert table.num_rows == 15
        assert table.column("a").to_pylist() == list(range(15))
        assert sum(yielded) <= 15  # 100 rows in the file, only a sliver was decoded

    def test_whole_file_and_small_files(self, parquet_path):
        inferrer = DataTypeInferrer()
        assert inferrer.read_parquet_sample(parquet_path, None).num_rows == 100
        assert inferrer.read_parquet_sample(parquet_path, 1000).num_rows == 100
        assert inferrer.read_parquet_sample(parquet_path, 1).num_rows == 1

    def test_empty_parquet(self, tmp_path):
        path = tmp_path / "empty.parquet"
        pq.write_table(pa.table({"a": pa.array([], pa.int64())}), path)
        table = DataTypeInferrer().read_parquet_sample(path, 10)
        assert table.num_rows == 0
        assert table.column_names == ["a"]

    def test_s3_forward_only_stream_is_buffered(self, parquet_path):
        inferrer = DataTypeInferrer()
        inferrer.io_handler = FakeS3(ForwardOnlyStream(parquet_path.read_bytes()))
        table = inferrer.read_parquet_sample("s3://bucket/groups.parquet", 12)
        assert inferrer.io_handler.calls == [("s3://bucket/groups.parquet", "binary")]
        assert table.num_rows == 12

    def test_types_are_preserved(self, tmp_path):
        path = tmp_path / "typed.parquet"
        source = pa.table(
            {
                "dec": pa.array([Decimal("1.2345")], pa.decimal128(18, 4)),
                "ts": pa.array([datetime.datetime(2020, 1, 1)], pa.timestamp("ns")),
            }
        )
        pq.write_table(source, path)
        table = DataTypeInferrer().read_parquet_sample(path, 10)
        assert table.schema.field("dec").type == pa.decimal128(18, 4)


class TestInputLocationGuard:
    @pytest.mark.parametrize(
        "location",
        [
            "http://example.com/a.csv",
            "https://example.com/a.csv",
            "ftp://example.com/a.csv",
            "file:///etc/passwd",
            "gs://bucket/a.csv",
            "s3a://bucket/a.csv",
            "sftp://host/a.csv",
        ],
    )
    @pytest.mark.parametrize(
        "reader", ["read_csv_sample", "read_excel_sample", "read_parquet_sample"]
    )
    def test_url_like_inputs_are_rejected(self, location, reader):
        with pytest.raises(ValueError, match="Unsupported input location"):
            getattr(DataTypeInferrer(), reader)(location, 10)

    def test_credentials_in_url_are_not_echoed(self):
        with pytest.raises(ValueError) as excinfo:
            DataTypeInferrer().read_csv_sample("https://user:hunter2@host/a.csv", 10)
        assert "hunter2" not in str(excinfo.value)

    def test_rejected_before_any_io(self, monkeypatch):
        def fail(*args, **kwargs):
            raise AssertionError("must not open anything")

        monkeypatch.setattr("builtins.open", fail)
        monkeypatch.setattr(inference_module.pv_csv, "open_csv", fail)
        with pytest.raises(ValueError):
            DataTypeInferrer().read_csv_sample("http://example.com/a.csv", 10)

    def test_generator_surfaces_the_guard(self):
        config = SchemaGenerationConfig(
            input_path="https://example.com/a.csv", file_type=FileType.CSV
        )
        with pytest.raises(ValueError, match="Unsupported input location"):
            SchemaGenerator(config).generate_schema()

    @pytest.mark.parametrize("path", ["C:\\data\\a.csv", "C:/data/a.csv", "c:\\a.csv"])
    def test_windows_drive_letters_are_not_urls(self, path):
        DataTypeInferrer._check_input_path(path)  # no ValueError
        with pytest.raises(OSError):  # simply a missing local file on this machine
            DataTypeInferrer().read_csv_sample(path, 10)

    def test_s3_uris_and_local_paths_pass(self, tmp_path):
        DataTypeInferrer._check_input_path("s3://bucket/key.csv")
        DataTypeInferrer._check_input_path(tmp_path / "x.csv")
        DataTypeInferrer._check_input_path("relative/dir/x.csv")


class TestSchemaGenerationEndToEnd:
    def test_csv_schema_types_do_not_depend_on_nrows(self, tmp_path):
        path = tmp_path / "e.csv"
        path.write_text(
            "id,zip,joined,score\n"
            + "".join(f"{i},0{i:04d},2021-01-{i % 28 + 1:02d},{i}.5\n" for i in range(1, 60))
        )
        limited = api.generate_schema_from_csv(path, nrows=10)
        unlimited = api.generate_schema_from_csv(path, nrows=None)
        assert limited["properties"] == unlimited["properties"]
        assert limited["properties"]["zip"] == {"type": "string"}
        assert limited["properties"]["id"] == {"type": "integer"}
        assert limited["properties"]["joined"] == {"type": "string", "format": "date"}
        assert limited["x-csv"]["dataTypes"]["joined"] == "date32"
        assert limited["x-generation"]["rows_analyzed"] == 10
        assert unlimited["x-generation"]["rows_analyzed"] == 59

    def test_zero_padded_codes_survive_into_the_schema(self, tmp_path):
        path = tmp_path / "z.csv"
        path.write_text("zip\n02134\n00123\n90210\n")
        schema = api.generate_schema_from_csv(path)
        assert schema["properties"]["zip"] == {"type": "string"}
        assert schema["x-csv"]["dataTypes"]["zip"] == "string"


# ---------------------------------------------------------------------------
# 5. Arrow -> Forklift type strings
# ---------------------------------------------------------------------------

FAITHFUL_TYPES = [
    (pa.int8(), "int8"),
    (pa.int64(), "int64"),
    (pa.uint32(), "uint32"),
    (pa.float32(), "float32"),
    (pa.float64(), "double"),
    (pa.bool_(), "bool"),
    (pa.string(), "string"),
    (pa.binary(), "binary"),
    (pa.date32(), "date32"),
    (pa.date64(), "date64"),
    (pa.decimal128(18, 4), "decimal128(18,4)"),
    (pa.decimal128(38, 0), "decimal128(38,0)"),
    (pa.decimal128(10, 2), "decimal128(10,2)"),
    (pa.timestamp("s"), "timestamp[s]"),
    (pa.timestamp("ms"), "timestamp[ms]"),
    (pa.timestamp("us"), "timestamp[us]"),
    (pa.timestamp("ns"), "timestamp[ns]"),
    (pa.timestamp("us", tz="UTC"), "timestamp[us, tz=UTC]"),
    (pa.timestamp("ms", tz="America/New_York"), "timestamp[ms, tz=America/New_York]"),
    (pa.duration("s"), "duration[s]"),
    (pa.duration("ms"), "duration[ms]"),
    (pa.duration("us"), "duration[us]"),
    (pa.duration("ns"), "duration[ns]"),
    (pa.list_(pa.int32()), "list<int32>"),
    (pa.list_(pa.decimal128(5, 2)), "list<decimal128(5,2)>"),
    (pa.list_(pa.list_(pa.string())), "list<list<string>>"),
    (pa.list_(pa.timestamp("ns", tz="UTC")), "list<timestamp[ns, tz=UTC]>"),
    (
        pa.dictionary(pa.int32(), pa.string()),
        "dictionary<values=string, indices=int32>",
    ),
    (pa.dictionary(pa.int8(), pa.string()), "dictionary<values=string, indices=int8>"),
]

# Arrow types without an importer-accepted spelling map to the closest accepted type
APPROXIMATE_TYPES = [
    (pa.float16(), "float32"),
    (pa.large_string(), "string"),
    (pa.large_binary(), "binary"),
    (pa.binary(16), "binary"),
    (pa.large_list(pa.int64()), "list<int64>"),
    (pa.list_(pa.int16(), 3), "list<int16>"),
    (pa.map_(pa.string(), pa.int32()), "list<struct>"),
    (pa.struct([("a", pa.int32())]), "struct"),
    (pa.time32("s"), "string"),
    (pa.time64("us"), "string"),
    (pa.decimal256(40, 2), "string"),
    (pa.null(), "string"),
]


class TestTypeMapping:
    @pytest.mark.parametrize("arrow_type,expected", FAITHFUL_TYPES + APPROXIMATE_TYPES)
    def test_mapping(self, arrow_type, expected):
        assert get_parquet_type_string(arrow_type) == expected

    @pytest.mark.parametrize("arrow_type,expected", FAITHFUL_TYPES)
    def test_round_trip(self, arrow_type, expected):
        parsed = parquet_type_string_to_arrow(get_parquet_type_string(arrow_type))
        assert parsed.equals(arrow_type) or str(parsed) == str(arrow_type)
        # and the string form is stable
        assert get_parquet_type_string(parsed) == expected

    @pytest.mark.parametrize("arrow_type,expected", FAITHFUL_TYPES + APPROXIMATE_TYPES)
    def test_every_generated_string_is_accepted_by_the_importers(self, arrow_type, expected):
        generated = get_parquet_type_string(arrow_type)
        assert ParquetTypeValidator.is_valid_parquet_type(generated)
        assert CsvSchemaImporter({}, validate=False)._is_valid_parquet_type(generated)
        assert ExcelSchemaImporter({}, validate=False)._is_valid_parquet_type(generated)

    def test_old_lossy_mappings_are_gone(self):
        assert get_parquet_type_string(pa.decimal128(18, 4)) != "string"
        assert get_parquet_type_string(pa.timestamp("ns")) != "timestamp[ms]"
        assert get_parquet_type_string(pa.timestamp("us", tz="UTC")) != "timestamp[ms]"
        assert get_parquet_type_string(pa.duration("ns")) != "duration[ms]"

    def test_unparseable_strings_raise(self):
        for text in ("", "int128", "list<>", "timestamp[xs]", "decimal128(a,b)"):
            with pytest.raises(ValueError):
                parquet_type_string_to_arrow(text)

    def test_generated_schema_with_rich_types_loads_in_the_importer(self, tmp_path):
        path = tmp_path / "rich.parquet"
        table = pa.table(
            {
                "amount": pa.array([Decimal("1.2345")], pa.decimal128(18, 4)),
                "created": pa.array(
                    [datetime.datetime(2020, 1, 1, tzinfo=datetime.timezone.utc)],
                    pa.timestamp("us", tz="UTC"),
                ),
                "elapsed": pa.array([datetime.timedelta(seconds=5)], pa.duration("ns")),
                "tags": pa.array([["a"]], pa.list_(pa.string())),
            }
        )
        pq.write_table(table, path)
        config = SchemaGenerationConfig(input_path=str(path), file_type=FileType.CSV)
        generator = SchemaGenerator(config)
        sample = DataTypeInferrer().read_parquet_sample(path, None)
        schema = generator._generate_schema_from_table(sample)
        data_types = schema["x-csv"]["dataTypes"]
        assert data_types == {
            "amount": "decimal128(18,4)",
            "created": "timestamp[us, tz=UTC]",
            "elapsed": "duration[ns]",
            "tags": "list<string>",
        }
        schema["x-csv"]["parquetTypeMapping"] = dict(data_types)
        importer = CsvSchemaImporter(schema)  # validates the type mapping
        assert importer.get_parquet_type_mapping()["amount"] == "decimal128(18,4)"
        metadata_types = {
            name: column["parquet_type"]
            for name, column in schema["x-metadata"]["column_metadata"].items()
        }
        assert metadata_types == data_types


# ---------------------------------------------------------------------------
# 6. Primary key inference
# ---------------------------------------------------------------------------


def infer_pk(table):
    return ConfigurationParser()._infer_primary_key_from_metadata(table)


class TestPrimaryKeyInference:
    def test_near_unique_column_is_not_enforced(self):
        # 19 distinct values out of 20 rows (95%): the old code accepted this with
        # enforceUniqueness=True
        values = list(range(19)) + [0]
        assert infer_pk(pa.table({"user_id": values})) is None

    def test_exactly_unique_column_is_inferred(self):
        result = infer_pk(pa.table({"user_id": list(range(20)), "other": [1] * 20}))
        assert result["columns"] == ["user_id"]
        assert result["enforceUniqueness"] is True
        assert result["inference_metadata"]["uniqueness_ratio"] == 1.0
        assert result["inference_metadata"]["rows_analyzed"] == 20

    def test_nulls_disqualify(self):
        assert infer_pk(pa.table({"user_id": pa.array([1, 2, None])})) is None
        assert infer_pk(pa.table({"user_id": pa.array([1.0, 2.0, float("nan")])})) is None

    @pytest.mark.parametrize(
        "name", ["width", "paid", "valid", "android", "monkey", "idle", "bid"]
    )
    def test_substring_matches_do_not_count(self, name):
        assert infer_pk(pa.table({name: list(range(10))})) is None

    @pytest.mark.parametrize(
        "name",
        [
            "id",
            "ID",
            "user_id",
            "id_code",
            "userId",
            "UserID",
            "order-id",
            "uuid",
            "row_guid",
            "pk",
            "primary_key",
        ],
    )
    def test_whole_token_names_count(self, name):
        result = infer_pk(pa.table({name: list(range(10))}))
        assert result is not None and result["columns"] == [name]

    def test_prefers_better_named_key(self):
        table = pa.table(
            {"width": list(range(10)), "row_uuid": list(range(10)), "id": list(range(10))}
        )
        result = infer_pk(table)
        assert result["columns"] == ["id"]  # an `id` token scores higher than `uuid`
        assert result["inference_metadata"]["alternative_candidates"] == ["row_uuid"]

    def test_split_name_tokens(self):
        assert split_name_tokens("userId") == ["user", "id"]
        assert split_name_tokens("user_id") == ["user", "id"]
        assert split_name_tokens("clientIPAddress") == ["client", "ip", "address"]
        assert split_name_tokens("HTTPServer2") == ["http", "server", "2"]
        assert split_name_tokens("width") == ["width"]
        assert split_name_tokens("") == []


# ---------------------------------------------------------------------------
# 7. Special type detection
# ---------------------------------------------------------------------------


class TestSpecialTypeDetection:
    @pytest.mark.parametrize(
        "name",
        [
            "description",
            "ship_date",
            "membership",
            "tip",
            "script",
            "ipsum",
            "hotel",
            "telephoto",
            "machine",
            "macro",
            "macaroni",
            "emailing",
            "zipper",
            "ssnake",
            "dip_switch",
        ],
    )
    def test_names_match_whole_tokens_only(self, name):
        assert SpecialTypeDetector._detect_from_column_name(name) is None
        assert SpecialTypeDetector.detect_special_type(name, ["plain text"]) is None

    @pytest.mark.parametrize(
        "name,expected",
        [
            ("ssn", "ssn"),
            ("employeeSSN", "ssn"),
            ("social_security_number", "ssn"),
            ("tel", "phone"),
            ("customerPhone", "phone"),
            ("ip", "ip_address"),
            ("clientIP", "ip_address"),
            ("IPAddress", "ip_address"),
            ("server_ip_addr", "ip_address"),
            ("mac", "mac_address"),
            ("deviceMAC", "mac_address"),
            ("e_mail", "email"),
            ("userEmail", "email"),
            ("zipCode", "zip_code"),
            ("POSTAL_CODE", "zip_code"),
        ],
    )
    def test_token_names_are_detected(self, name, expected):
        assert SpecialTypeDetector._detect_from_column_name(name) == expected

    def test_nine_digit_ids_are_not_ssns(self):
        ids = ["123456789", "987654321", "555443333"]
        assert SpecialTypeDetector.detect_special_type("customer_ref", ids) is None
        assert SpecialTypeDetector.detect_special_type("ssn", ids) == "ssn"  # name still counts

    def test_ten_digit_epochs_are_not_phones(self):
        epochs = ["1700000000", "1700000060", "1700000120"]
        assert SpecialTypeDetector.detect_special_type("created", epochs) is None

    def test_five_digit_numbers_are_not_zip_codes(self):
        counts = ["12345", "54321", "10001"]
        assert SpecialTypeDetector.detect_special_type("order_count", counts) is None
        assert SpecialTypeDetector.detect_special_type("zip", counts) == "zip_code"

    def test_separated_formats_are_detected(self):
        assert SpecialTypeDetector.detect_special_type("c", ["123-45-6789"] * 3) == "ssn"
        assert SpecialTypeDetector.detect_special_type("c", ["(123) 456-7890"] * 3) == "phone"
        assert SpecialTypeDetector.detect_special_type("c", ["123.456.7890"] * 3) == "phone"
        assert SpecialTypeDetector.detect_special_type("c", ["+1 123-456-7890"] * 3) == "phone"
        assert SpecialTypeDetector.detect_special_type("c", ["12345-6789"] * 3) == "zip_code"

    def test_content_must_match_the_whole_value(self):
        text = ["call 123-456-7890 now", "ssn is 123-45-6789 ok", "mail me a@b.com please"]
        assert SpecialTypeDetector.detect_special_type("notes", text) is None

    def test_ip_validation(self):
        assert (
            SpecialTypeDetector.detect_special_type("c", ["10.0.0.1", "192.168.0.255"])
            == "ip_address"
        )
        assert SpecialTypeDetector.detect_special_type("c", ["::1", "fe80::1"]) == "ip_address"
        assert SpecialTypeDetector.detect_special_type("c", ["999.1.1.1", "300.2.2.2"]) is None
        assert SpecialTypeDetector.detect_special_type("c", ["1.2.3", "4.5.6"]) is None

    def test_email_pattern_has_no_stray_pipe_in_tld(self):
        assert SpecialTypeDetector.detect_special_type("c", ["a@b.c|"]) is None
        assert SpecialTypeDetector.detect_special_type("c", ["a@b.co", "x@y.org"]) == "email"

    def test_mac_formats(self):
        assert SpecialTypeDetector.detect_special_type("c", ["aabb.ccdd.eeff"]) == "mac_address"
        assert SpecialTypeDetector.detect_special_type("c", ["AA-BB-CC-DD-EE-FF"]) == "mac_address"

    def test_zero_like_values_are_still_values(self):
        assert SpecialTypeDetector._detect_from_content([0, 0, 0], 0.7) is None
        assert SpecialTypeDetector._detect_from_content(["", None, " "], 0.7) is None
