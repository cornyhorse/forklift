"""Data type inference from file contents.

Samples are read with PyArrow only (no pandas):

* CSV files are streamed with ``pyarrow.csv.open_csv`` **as all-string columns** and the
  stream is closed as soon as ``nrows`` rows have been read, so large local or S3 files are
  never read completely. Forklift's own type inference then runs over the strings, so the
  result does not depend on how many rows were sampled and identifiers such as ``02134`` keep
  their leading zeros.
* Excel files are read with ``openpyxl`` in read-only mode, limited to ``nrows + 1`` rows.
* Parquet files are read row group by row group until ``nrows`` rows are available.

Only local paths and ``s3://`` URIs are accepted; other URL-like inputs (``http://``,
``ftp://``, ``file://`` ...) are rejected so a schema request cannot be used to make the
process fetch arbitrary locations.
"""

import codecs
import csv
import io
import numbers
import re
from pathlib import Path
from typing import Any, Callable, ContextManager, Iterable, List, Optional, Union

import pyarrow as pa
import pyarrow.compute as pc
import pyarrow.csv as pv_csv
import pyarrow.parquet as pq

from ...io import UnifiedIOHandler, is_s3_path

# Strings that are read as null when inferring types from CSV text, in addition to the empty
# string. "NA" is deliberately absent: it is a legitimate value (country code of Namibia,
# "North America", a sodium symbol ...) and silently turning it into null loses data.
DEFAULT_NULL_TOKENS = ("", "NULL", "null", "N/A", "n/a", "#N/A", "NaN", "nan")

# Integers: no sign other than '-', no leading zeros (identifiers like 02134 stay strings).
_INTEGER_PATTERN = r"^-?(0|[1-9][0-9]*)$"
# Plain decimal/scientific numbers with the same leading-zero rule for the integer part.
_NUMBER_PATTERN = r"^-?((0|[1-9][0-9]*)(\.[0-9]*)?|\.[0-9]+)([eE][+-]?[0-9]+)?$"
_DATE_PATTERN = r"^[0-9]{4}-[0-9]{2}-[0-9]{2}$"
_TIMESTAMP_PATTERN = (
    r"^[0-9]{4}-[0-9]{2}-[0-9]{2}[T ][0-9]{2}:[0-9]{2}(:[0-9]{2}(\.[0-9]{1,9})?)?"
    r"(Z|[+-][0-9]{2}:?[0-9]{2})?$"
)
_TIMEZONE_SUFFIX_PATTERN = r"(Z|[+-][0-9]{2}:?[0-9]{2})$"

_URL_LIKE = re.compile(r"^[A-Za-z][A-Za-z0-9+.\-]+://")

# CSV block sizes tried in order (see ``_read_string_csv``)
_BLOCK_SIZES = (1 << 16, 1 << 20, 1 << 24, 1 << 26)

# ``default_column_type`` (all columns as strings in one pass) exists in recent pyarrow only.
_SUPPORTS_DEFAULT_COLUMN_TYPE = hasattr(pv_csv.ConvertOptions, "default_column_type")


class DataTypeInferrer:
    """Handles data type inference from various file formats."""

    NULL_TOKENS = DEFAULT_NULL_TOKENS

    def __init__(self):
        self.io_handler = UnifiedIOHandler()

    # ------------------------------------------------------------------
    # CSV
    # ------------------------------------------------------------------

    def read_csv_sample(
        self,
        input_path: Union[str, Path],
        nrows: Optional[int] = 1000,
        delimiter: str = ",",
        encoding: str = "utf-8",
    ) -> pa.Table:
        """Read CSV sample data for inference.

        The file (local or ``s3://``) is streamed and reading stops after ``nrows`` data rows.
        Column types are inferred from the text values (see :meth:`infer_types_from_strings`),
        so ``nrows=None`` and a sufficiently large ``nrows`` give the same result.

        Args:
            input_path: Path to CSV file
            nrows: Number of rows to read for analysis (``None`` reads the whole file)
            delimiter: CSV delimiter
            encoding: File encoding

        Returns:
            pa.Table: Sample data as PyArrow table

        Raises:
            ValueError: For unsupported (URL-like) inputs, an invalid ``nrows`` or encoding,
                or unreadable CSV data. Messages never contain cell values.
        """
        self._check_input_path(input_path)
        nrows = self._check_nrows(nrows)
        try:
            codecs.lookup(encoding)
        except LookupError:
            raise ValueError(f"Unknown encoding: {encoding!r}") from None

        raw = self._read_string_csv(
            self._binary_opener(input_path, seekable=False), nrows, delimiter, encoding
        )
        return self.infer_types_from_strings(raw)

    def _read_string_csv(
        self,
        opener: Callable[[], ContextManager],
        nrows: Optional[int],
        delimiter: str,
        encoding: str,
    ) -> pa.Table:
        """Stream the first ``nrows`` CSV rows as an all-string table.

        Arrow prefetches roughly 32 blocks ahead of the consumer, so a sample is read with a
        small block size first (a few MiB of I/O for typical rows). Records that do not fit
        into a block make Arrow fail with a "straddles" error; the read is then repeated with
        the next larger block size.
        """
        # A full read gains nothing from small blocks
        block_sizes = _BLOCK_SIZES if nrows is not None else _BLOCK_SIZES[1:]
        parse_options = pv_csv.ParseOptions(
            delimiter=delimiter, quote_char='"', double_quote=True, newlines_in_values=True
        )

        for attempt, block_size in enumerate(block_sizes):
            is_last = attempt == len(block_sizes) - 1
            read_options = pv_csv.ReadOptions(
                encoding=encoding, use_threads=False, block_size=block_size
            )
            try:
                table = self._stream_string_csv(opener, read_options, parse_options, nrows)
                break
            except pa.ArrowInvalid as exc:
                message = str(exc)
                if "cannot infer number of columns" in message:
                    # Either a header line without a trailing newline (Arrow needs a row
                    # terminator) or a header longer than the block
                    table = self._header_only_table(opener, delimiter, encoding, block_size)
                    if table is not None:
                        break
                elif "straddles" not in message:
                    raise ValueError(
                        f"Failed to read CSV sample: {self._csv_error(exc)}"
                    ) from None
                if is_last:
                    raise ValueError(
                        "Failed to read CSV sample: a header or record is larger than the "
                        "maximum supported size"
                    ) from None
            except UnicodeError:
                raise ValueError(
                    f"Failed to read CSV sample: data is not valid {encoding} text"
                ) from None

        if nrows is not None and table.num_rows > nrows:
            table = table.slice(0, nrows)
        return table.rename_columns(self._clean_column_names(table.column_names))

    def _stream_string_csv(
        self,
        opener: Callable[[], ContextManager],
        read_options: pv_csv.ReadOptions,
        parse_options: pv_csv.ParseOptions,
        nrows: Optional[int],
    ) -> pa.Table:
        convert_options = self._string_convert_options(opener, read_options, parse_options)
        with opener() as stream:
            reader = pv_csv.open_csv(
                stream,
                read_options=read_options,
                parse_options=parse_options,
                convert_options=convert_options,
            )
            try:
                batches = []
                rows = 0
                for batch in reader:
                    batches.append(batch)
                    rows += batch.num_rows
                    if nrows is not None and rows >= nrows:
                        break
                return pa.Table.from_batches(batches, schema=reader.schema)
            finally:
                reader.close()

    @staticmethod
    def _string_convert_options(
        opener: Callable[[], ContextManager],
        read_options: pv_csv.ReadOptions,
        parse_options: pv_csv.ParseOptions,
    ) -> pv_csv.ConvertOptions:
        """Convert options that keep every column as an unmodified string."""
        if _SUPPORTS_DEFAULT_COLUMN_TYPE:
            return pv_csv.ConvertOptions(default_column_type=pa.string())

        # Older pyarrow: the header has to be known to pin every column to string
        with opener() as stream:
            probe = pv_csv.open_csv(stream, read_options=read_options, parse_options=parse_options)
            try:
                names = probe.schema.names
            finally:
                probe.close()
        return pv_csv.ConvertOptions(column_types={name: pa.string() for name in names})

    @staticmethod
    def _header_only_table(
        opener: Callable[[], ContextManager], delimiter: str, encoding: str, block_size: int
    ) -> Optional[pa.Table]:
        """Empty all-string table if the file consists of a header line only.

        Returns ``None`` when the file is longer than ``block_size`` (so Arrow's failure was
        about the block size, not about a missing row terminator).
        """
        with opener() as stream:
            head = stream.read(block_size)
            if len(head) >= block_size or stream.read(1):
                return None
        try:
            text = head.decode(encoding)
            header = next(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter), [])
        except (UnicodeError, csv.Error):
            header = []
        if not header:
            raise ValueError("Failed to read CSV sample: file is empty") from None
        return pa.table({name: pa.array([], pa.string()) for name in header})

    @staticmethod
    def _csv_error(exc: Exception) -> str:
        """Describe an Arrow CSV error without echoing the offending data.

        Arrow quotes the offending row in its parse errors; only positions and counts are kept.
        """
        message = str(exc)
        match = re.search(r"Row #(\d+): Expected (\d+) columns, got (\d+)", message)
        if match:
            row, expected, got = match.groups()
            return f"row {row} has {got} columns, expected {expected}"
        if "Empty CSV file" in message:
            return "file is empty"
        if "invalid UTF8" in message:
            row = re.search(r"Row #(\d+)", message)
            where = f" (row {row.group(1)})" if row else ""
            return f"invalid text for the configured encoding{where}"
        return "invalid CSV data"

    # ------------------------------------------------------------------
    # Type inference over strings
    # ------------------------------------------------------------------

    def infer_types_from_strings(
        self, table: pa.Table, null_tokens: Optional[Iterable[str]] = None
    ) -> pa.Table:
        """Infer column types for a table of string columns.

        Per column, after turning the empty string and ``null_tokens`` (default
        :data:`DEFAULT_NULL_TOKENS`) into nulls, the first matching rule wins:

        1. ``boolean``: every value is ``true``/``false`` (any case)
        2. ``int64`` (or ``uint64``): every value matches ``^-?(0|[1-9]\\d*)$``. Leading zeros
           are therefore never integers; values outside the 64-bit range stay strings
        3. ``float64``: every value is a plain decimal / scientific number
        4. ``date32``: every value is a valid ``YYYY-MM-DD`` date
        5. ``timestamp``: every value is ``YYYY-MM-DD[T ]HH:MM[:SS[.f]]`` with or without a UTC
           offset (all values must agree); the unit is as fine as the fractions require
        6. otherwise ``string``

        Args:
            table: Table whose columns are string arrays
            null_tokens: Strings to read as null besides the empty string

        Returns:
            pa.Table: Table with inferred column types
        """
        tokens = list(self.NULL_TOKENS if null_tokens is None else null_tokens)
        if "" not in tokens:
            tokens.append("")
        token_set = pa.array(tokens, pa.string())

        columns = []
        for column in table.columns:
            if not (pa.types.is_string(column.type) or pa.types.is_large_string(column.type)):
                columns.append(column)
                continue
            is_null_token = pc.is_in(column, value_set=token_set)
            nulled = pc.if_else(is_null_token, pa.scalar(None, column.type), column)
            columns.append(self._infer_string_column(nulled))
        return pa.Table.from_arrays(columns, names=table.column_names)

    def _infer_string_column(self, column: pa.ChunkedArray) -> pa.ChunkedArray:
        """Cast one string column (nulls already applied) to the narrowest matching type."""
        values = column.drop_null()
        if len(values) == 0:
            return column

        try:
            if self._all_match_values(values, ["true", "false"], lowercase=True):
                return pc.equal(pc.utf8_lower(column), "true")

            if self._all_match(values, _INTEGER_PATTERN):
                return self._cast_integer(column, values)

            if self._all_match(values, _NUMBER_PATTERN):
                return column.cast(pa.float64())

            if self._all_match(values, _DATE_PATTERN):
                return column.cast(pa.date32())

            if self._all_match(values, _TIMESTAMP_PATTERN):
                return self._cast_timestamp(column, values)
        except pa.ArrowInvalid:
            # e.g. 2023-02-30 or an out-of-range number: it is not that type after all
            pass

        return column

    @staticmethod
    def _all_match(values: pa.ChunkedArray, pattern: str) -> bool:
        return bool(pc.all(pc.match_substring_regex(values, pattern)).as_py())

    @staticmethod
    def _all_match_values(values: pa.ChunkedArray, allowed: List[str], lowercase: bool) -> bool:
        candidates = pc.utf8_lower(values) if lowercase else values
        return bool(pc.all(pc.is_in(candidates, value_set=pa.array(allowed))).as_py())

    @staticmethod
    def _cast_integer(column: pa.ChunkedArray, values: pa.ChunkedArray) -> pa.ChunkedArray:
        try:
            return column.cast(pa.int64())
        except pa.ArrowInvalid:
            pass
        has_negative = bool(pc.any(pc.starts_with(values, "-")).as_py())
        if not has_negative:
            try:
                return column.cast(pa.uint64())
            except pa.ArrowInvalid:
                pass
        # Beyond 64 bits: keep the digits as text instead of losing precision
        return column

    def _cast_timestamp(self, column: pa.ChunkedArray, values: pa.ChunkedArray) -> pa.ChunkedArray:
        zoned = pc.match_substring_regex(values, _TIMEZONE_SUFFIX_PATTERN)
        if bool(pc.all(zoned).as_py()):
            tz = "UTC"
        elif not bool(pc.any(zoned).as_py()):
            tz = None
        else:
            return column  # mix of naive and zone-aware timestamps

        if bool(pc.any(pc.match_substring_regex(values, r"\.[0-9]{7,9}")).as_py()):
            unit = "ns"
        elif bool(pc.any(pc.match_substring_regex(values, r"\.[0-9]{4,6}")).as_py()):
            unit = "us"
        elif bool(pc.any(pc.match_substring_regex(values, r"\.[0-9]{1,3}")).as_py()):
            unit = "ms"
        else:
            unit = "s"
        return column.cast(pa.timestamp(unit, tz=tz))

    # ------------------------------------------------------------------
    # Excel
    # ------------------------------------------------------------------

    def read_excel_sample(
        self,
        input_path: Union[str, Path],
        nrows: Optional[int] = 1000,
        sheet_name: Union[str, int, None] = None,
    ) -> pa.Table:
        """Read Excel sample data for inference.

        The workbook is opened read-only with ``openpyxl`` and only the header row plus
        ``nrows`` data rows are read. Cell values keep their Excel types.

        Args:
            input_path: Path to Excel file (``.xlsx`` / ``.xlsm``; local or ``s3://``)
            nrows: Number of rows to read for analysis (``None`` reads the whole sheet)
            sheet_name: Sheet name or index (default: first sheet)

        Returns:
            pa.Table: Sample data as PyArrow table
        """
        self._check_input_path(input_path)
        nrows = self._check_nrows(nrows)
        if str(input_path).lower().endswith(".xls"):
            raise ValueError(
                "Legacy .xls workbooks are not supported for schema generation; "
                "convert the file to .xlsx"
            )

        import openpyxl

        with self._binary_opener(input_path)() as source:
            workbook = openpyxl.load_workbook(
                self._seekable(source), read_only=True, data_only=True
            )
            try:
                worksheet = self._select_worksheet(workbook, sheet_name)
                # Some writers store a wrong <dimension>; openpyxl would then drop columns/rows
                worksheet.reset_dimensions()
                max_row = None if nrows is None else nrows + 1
                rows = list(worksheet.iter_rows(values_only=True, max_row=max_row))
            finally:
                workbook.close()

        return self._table_from_rows(rows)

    @staticmethod
    def _select_worksheet(workbook, sheet_name: Union[str, int, None]):
        if sheet_name is None:
            sheet_name = 0
        if isinstance(sheet_name, numbers.Integral) and not isinstance(sheet_name, bool):
            try:
                return workbook.worksheets[int(sheet_name)]
            except IndexError:
                raise ValueError(f"Sheet index {sheet_name} is out of range") from None
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet_name!r} not found in workbook")
        return workbook[sheet_name]

    def _table_from_rows(self, rows: List[tuple]) -> pa.Table:
        """Build an Arrow table from Excel rows (first row is the header)."""
        # Rows that are entirely empty at the end are formatting leftovers, not data
        while rows and all(cell is None for cell in rows[-1]):
            rows.pop()
        if not rows:
            return pa.table({})

        header, data = rows[0], rows[1:]
        width = max([len(header)] + [len(row) for row in data])
        names = [
            str(header[i]) if i < len(header) and header[i] is not None else ""
            for i in range(width)
        ]
        arrays = []
        for i in range(width):
            values = [row[i] if i < len(row) else None for row in data]
            arrays.append(self._array_from_values(values))
        return pa.Table.from_arrays(arrays, names=self._clean_column_names(names))

    @staticmethod
    def _array_from_values(values: list) -> pa.Array:
        """Arrow array for a column of Python values; mixed types fall back to strings."""
        try:
            array = pa.array(values)
        except (pa.ArrowInvalid, pa.ArrowTypeError, OverflowError):
            return pa.array([None if v is None else str(v) for v in values], pa.string())
        if pa.types.is_null(array.type):
            return array.cast(pa.string())
        return array

    # ------------------------------------------------------------------
    # Parquet
    # ------------------------------------------------------------------

    def read_parquet_sample(
        self, input_path: Union[str, Path], nrows: Optional[int] = 1000
    ) -> pa.Table:
        """Read Parquet sample data for inference.

        Only as many row groups as needed for ``nrows`` rows are read.

        Args:
            input_path: Path to Parquet file (local or ``s3://``)
            nrows: Number of rows to read for analysis (``None`` reads the whole file)

        Returns:
            pa.Table: Sample data as PyArrow table
        """
        self._check_input_path(input_path)
        nrows = self._check_nrows(nrows)

        if is_s3_path(str(input_path)):
            with self._binary_opener(input_path)() as source:
                return self._sample_parquet(pq.ParquetFile(self._seekable(source)), nrows)

        parquet_file = pq.ParquetFile(str(input_path))
        try:
            return self._sample_parquet(parquet_file, nrows)
        finally:
            close = getattr(parquet_file, "close", None)
            if callable(close):
                close()

    @staticmethod
    def _sample_parquet(parquet_file: "pq.ParquetFile", nrows: Optional[int]) -> pa.Table:
        if nrows is None:
            return parquet_file.read()

        batches = []
        rows = 0
        for batch in parquet_file.iter_batches(batch_size=min(nrows, 65536)):
            batches.append(batch)
            rows += batch.num_rows
            if rows >= nrows:
                break
        table = pa.Table.from_batches(batches, schema=parquet_file.schema_arrow)
        return table.slice(0, nrows)

    # ------------------------------------------------------------------
    # Shared helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _check_input_path(input_path: Union[str, Path]) -> None:
        """Allow local paths and ``s3://`` URIs only (guards against SSRF-style inputs)."""
        text = str(input_path)
        if is_s3_path(text):
            return
        if _URL_LIKE.match(text):
            scheme = text.split("://", 1)[0]
            raise ValueError(
                f"Unsupported input location (scheme {scheme!r}): only local file paths "
                "and s3:// URIs can be analysed"
            )

    @staticmethod
    def _check_nrows(nrows: Optional[int]) -> Optional[int]:
        if nrows is None:
            return None
        if isinstance(nrows, bool) or not isinstance(nrows, numbers.Integral) or nrows < 1:
            raise ValueError("nrows must be a positive integer or None")
        return int(nrows)

    def _binary_opener(
        self, input_path: Union[str, Path], seekable: bool = True
    ) -> Callable[[], ContextManager]:
        """Factory for binary file objects of a local file or an S3 object.

        ``seekable=False`` asks the IO layer for the raw forward-only S3 stream instead of a
        full local copy, so the CSV sampler can stop reading once it has ``nrows`` records.
        """
        if is_s3_path(str(input_path)):
            return lambda: self.io_handler.open_for_read(
                str(input_path), encoding="binary", seekable=seekable
            )
        return lambda: open(input_path, "rb")

    @staticmethod
    def _seekable(stream: Any) -> Any:
        """Return ``stream`` itself if it supports random access, else an in-memory copy.

        Parquet and Excel readers need random access; some remote streams only read forward.
        """
        seekable = getattr(stream, "seekable", None)
        if callable(seekable) and seekable():
            return stream
        return io.BytesIO(stream.read())

    @staticmethod
    def _clean_column_names(names: Iterable[str]) -> List[str]:
        """Give empty names a placeholder and make duplicates unique (``a``, ``a.1``)."""
        used = set()
        cleaned = []
        for position, raw in enumerate(names):
            name = raw if raw not in (None, "") else f"Unnamed: {position}"
            candidate, suffix = name, 0
            while candidate in used:
                suffix += 1
                candidate = f"{name}.{suffix}"
            used.add(candidate)
            cleaned.append(candidate)
        return cleaned

    def infer_schema_from_data(self, table: pa.Table) -> dict:
        """Generate basic schema structure from table data.

        Args:
            table: PyArrow table to analyze

        Returns:
            dict: Basic schema structure with inferred types
        """
        from ..types.data_types import DataTypeConverter

        converter = DataTypeConverter()
        properties = {}

        for field in table.schema:
            column_name = field.name
            arrow_type = field.type

            # Convert Arrow type to JSON Schema type
            json_type = converter.arrow_to_json_schema_type(arrow_type)
            properties[column_name] = json_type

        return {"type": "object", "properties": properties, "additionalProperties": False}
